# -*- coding: utf-8 -*-
# NOTE(LP-343): 移植自交付版 generate_final_report_docx.py，勿直接改业务逻辑
"""根据步骤1-5结构化判别结果生成正式Word报告。"""

from __future__ import annotations

import argparse
import json
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE = "2E74B5"
DARK = "1F4D78"
PALE = "E8EEF5"
LIGHT = "F4F6F9"
GOLD = "FFF2CC"


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = SCRIPT_DIR.parent / "templates" / "综合判别报告模板.docx"


def has_value(value: Any) -> bool:
    return value is not None and str(value).strip() not in {"", "None", "nan", "NaN"}


def safe_get(data: Mapping[str, Any], *keys: str, default: Any = None) -> Any:
    value: Any = data
    for key in keys:
        if not isinstance(value, Mapping) or key not in value:
            return default
        value = value[key]
    return value if has_value(value) else default


def display(value: Any, empty: str = "未提供") -> str:
    if not has_value(value):
        return empty
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def format_number(value: Any, digits: int = 3, empty: str = "未提供") -> str:
    if not has_value(value):
        return empty
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    text = f"{number:.{digits}f}".rstrip("0").rstrip(".")
    return "0" if text == "-0" else text


def value_with_unit(value: Any, unit: str, digits: int = 3) -> str:
    return f"{format_number(value, digits)} {unit}" if has_value(value) else "未提供"


def report_identifier(result: Mapping[str, Any]) -> str:
    event_id = safe_get(result, "inputs", "lightning", "event_id")
    if has_value(event_id):
        return f"FD-{event_id}"
    fault_time = str(safe_get(result, "inputs", "fault_record", "fault_time", default=""))
    compact = "".join(ch for ch in fault_time if ch.isdigit())[:14]
    return f"FD-{compact}" if compact else "FD-未编号"


def normalize_phase(value: Any) -> str:
    return "".join(dict.fromkeys(ch for ch in str(value or "").upper() if ch in "ABC"))


def normalize_fault_phase_code(value: Any) -> str:
    """报告统一使用N表示接地标志，同时兼容读取历史G/中文接地写法。"""
    text = str(value or "").strip().upper()
    phase = normalize_phase(text)
    has_ground_marker = "N" in text or "G" in text or "接地" in text or "GROUND" in text
    return f"{phase}N" if phase and has_ground_marker else phase


def sequence_phase(value: Any) -> str:
    if isinstance(value, Sequence) and not isinstance(value, str):
        value = "".join(str(item) for item in value)
    return normalize_phase(value)


def phases_have_containment(first: Any, second: Any) -> bool:
    first_phase = set(sequence_phase(first))
    second_phase = set(sequence_phase(second))
    return bool(first_phase and second_phase and (first_phase <= second_phase or second_phase <= first_phase))


def wave_analysis_phase(wave: Mapping[str, Any]) -> str:
    """行波结论使用行波自身相别；发生规则修正时使用修正采用的相别。"""
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    if model.get("correction_applied"):
        corrected = sequence_phase(model.get("corrected_fault_phase") or wave.get("recording_fault_phase"))
        if corrected:
            return corrected
    main_phase = sequence_phase(model.get("main_phase") or review.get("initial_wave_main_phase"))
    if main_phase:
        return main_phase
    return sequence_phase(
        review.get("initial_wave_active_phases")
        or model.get("active_phases")
        or wave.get("wave_active_phases")
    )


def wave_candidate_phase(wave: Mapping[str, Any]) -> str:
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    return sequence_phase(
        review.get("initial_wave_active_phases")
        or model.get("active_phases")
        or wave.get("wave_active_phases")
    )


def corrected_fault_phase(result: Mapping[str, Any]) -> str:
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    return sequence_phase(
        safe_get(result, "inputs", "fault_record", "fault_phase")
        or wave.get("recording_fault_phase")
        or model.get("corrected_fault_phase")
    )


def wave_conclusion_phase(result: Mapping[str, Any]) -> str:
    """行波章节采用的相别。

    反击结论只在录波故障相与行波候补故障相一致时提前采用录波相别；否则保留行波自身故障相。
    """
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    label = str(model.get("final_label") or wave.get("classification") or "").upper()
    classification = str(wave.get("classification") or "")
    if label == "F" or "反击" in classification:
        recording_phase = corrected_fault_phase(result)
        candidate_phase = wave_candidate_phase(wave)
        if recording_phase and candidate_phase and recording_phase == candidate_phase:
            return recording_phase
        review = wave.get("fault_phase_review", {})
        wave_phase = sequence_phase(model.get("main_phase") or review.get("initial_wave_main_phase"))
        return wave_phase or candidate_phase
    return wave_analysis_phase(wave)


def phase_conductor_text(phase: str) -> str:
    if len(phase) == 2:
        return f"{phase}两相导线"
    if len(phase) == 3:
        return f"{phase}三相导线"
    return f"{phase}相导线" if phase else "相关导线"


def lightning_fault_text(label: Any, subtype: Any, phase: Any, *, detailed_subtype: bool) -> str:
    """按判别类别和相别生成完整故障原因句，不直接复用模型的简写标签。"""
    label_text = str(label or "").upper()
    subtype_text_value = str(subtype or "").upper()
    # 部分MATLAB版本会直接把绕击亚型写入final_label，而不另写R。
    if label_text in {"SF", "FBF", "FSF", "SF_FBF_REVIEW"} and not subtype_text_value:
        subtype_text_value = label_text
        label_text = "R"
    phase_text = sequence_phase(phase)
    if label_text == "F" or "反击" in str(label):
        return f"雷电反击导致{phase_text}相跳闸故障" if phase_text else "雷电反击故障"
    if label_text == "R" or "绕击" in str(label):
        conductor = phase_conductor_text(phase_text)
        if not detailed_subtype:
            return f"雷电绕击{conductor}跳闸故障"
        if subtype_text_value == "SF":
            return f"普通绕击，可能是雷电绕击{conductor}跳闸故障"
        if subtype_text_value == "FBF":
            return f"雷击避雷线与导线型绕击，可能是雷电绕击避雷线与{conductor}故障"
        if subtype_text_value == "FSF":
            return f"雷击多相导线型绕击，可能是雷电绕击{conductor}故障"
        if subtype_text_value == "SF_FBF_REVIEW":
            return (
                f"可能是雷电绕击{conductor}故障；故障类型处于普通绕击与雷击避雷线和导线型绕击边界，"
                "行波特征倾向后者，建议人工复核"
            )
        return f"雷电绕击{conductor}故障（故障类型不能唯一确定，建议人工复核）"
    if label_text == "N" or "非雷击" in str(label):
        return f"非雷击导致{phase_text}相跳闸故障" if phase_text else "非雷击故障"
    return "故障原因待复核"


def wave_fault_conclusion(result: Mapping[str, Any]) -> str:
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    label = model.get("final_label") or wave.get("classification")
    return lightning_fault_text(
        label,
        model.get("final_subtype"),
        wave_conclusion_phase(result),
        detailed_subtype=True,
    )


def ordinary_shielding_multi_phase_conflict(result: Mapping[str, Any]) -> bool:
    """普通绕击只用于单相故障；录波为多相时，融合类型应提升为多相导线型绕击。"""
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    label = str(model.get("final_label") or wave.get("classification") or "").upper()
    subtype = str(model.get("final_subtype") or "").upper()
    return (label == "R" or "绕击" in str(wave.get("classification") or "")) and subtype == "SF" and len(
        corrected_fault_phase(result)
    ) >= 2


def fused_fault_subtype(result: Mapping[str, Any]) -> str:
    model = safe_get(result, "traveling_wave_analysis", "model_result", default={}) or {}
    if ordinary_shielding_multi_phase_conflict(result):
        return "FSF"
    subtype = str(model.get("final_subtype") or "").upper()
    if subtype:
        return subtype
    direct_label = str(model.get("final_label") or "").upper()
    return direct_label if direct_label in {"SF", "FBF", "FSF", "SF_FBF_REVIEW"} else ""


def fused_fault_conclusion(result: Mapping[str, Any]) -> str:
    fusion = result.get("cause_fusion", {})
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    mechanism = fusion.get("final_mechanism") or model.get("final_label")
    subtype = fused_fault_subtype(result)
    # 交付版可能接收不同版本的MATLAB字段。融合对象没有写机制时，依次从
    # 行波最终标签、亚型、分类中文结论和融合详情恢复，避免已通过结果落成“待复核”。
    evidence_text = " ".join(
        str(item or "")
        for item in (
            mechanism,
            model.get("final_label"),
            subtype,
            wave.get("classification"),
            wave.get("initial_classification"),
            fusion.get("final_detail"),
        )
    ).upper()
    if not mechanism or str(mechanism).upper() in {"L", "雷击", "LIGHTNING"}:
        if subtype in {"SF", "FBF", "FSF", "SF_FBF_REVIEW"} or any(
            token in evidence_text for token in ("绕击", "SHIELDING", "SF", "FBF", "FSF")
        ):
            mechanism = "R"
        elif "反击" in evidence_text or "BACKFLASH" in evidence_text:
            mechanism = "F"
        elif "非雷击" in evidence_text or re.search(r"(^|\s)N($|\s)", evidence_text):
            mechanism = "N"
    return lightning_fault_text(
        mechanism,
        subtype,
        corrected_fault_phase(result),
        detailed_subtype=True,
    )


def effective_review_decision(result: Mapping[str, Any]) -> str:
    if ordinary_shielding_multi_phase_conflict(result):
        return "建议人工复核"
    return display(safe_get(result, "overall_recommendation", "decision"), "建议人工复核")


def effective_review_summary(result: Mapping[str, Any]) -> str:
    if ordinary_shielding_multi_phase_conflict(result):
        return "录波检测为多相故障，但行波仅符合普通绕击特征；融合类型按雷击多相导线型绕击处理，行波特征尚未满足对应判据，建议人工复核"
    return display(safe_get(result, "overall_recommendation", "summary"), "现有证据不足")


def wave_feature_conclusion(result: Mapping[str, Any]) -> str:
    """以中文全称概括行波最终采用的特征类型。"""
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    label = str(model.get("final_label") or "").upper()
    subtype = str(model.get("final_subtype") or "").upper()
    if label == "N":
        return "非雷击特征"
    if label == "F":
        return "雷电反击特征"
    if label == "R":
        if subtype == "SF_FBF_REVIEW":
            return "普通绕击与雷击避雷线和导线型绕击边界特征（倾向后者）"
        return f"{subtype_text(subtype)}特征" if subtype in {"SF", "FBF", "FSF"} else "雷电绕击特征"
    classification = str(wave.get("classification") or "")
    if "非雷击" in classification:
        return "非雷击特征"
    if "反击" in classification:
        return "雷电反击特征"
    if "绕击" in classification:
        return "雷电绕击特征"
    return "待复核的行波特征"


def fusion_conclusion_sentence(result: Mapping[str, Any]) -> str:
    """按幅值、录波相别、行波特征和综合结论四项生成完整融合语句。"""
    lightning_cause = result.get("lightning_cause_analysis", {})
    amplitude = display(
        lightning_cause.get("report_text")
        or lightning_cause.get("conclusion")
        or lightning_cause.get("mechanism"),
        "耐雷水平数据不足，暂无法由雷电流幅值判定故障机制",
    ).rstrip("。；")
    phase = corrected_fault_phase(result) or "未提供"
    return (
        f"根据雷电流幅值判断，{user_facing_model_terms(amplitude)}；"
        f"录波故障相为{phase}；"
        f"行波特征符合{wave_feature_conclusion(result)}；"
        f"综合判断为{fused_fault_conclusion(result)}。"
    )


def status_display(value: Any) -> str:
    return {
        "ok": "有效",
        "decided": "已判定",
        "needs_review": "需人工复核",
        "unavailable": "无法计算",
        "not_applicable": "不适用",
        "missing": "未提供",
    }.get(str(value or "").strip().lower(), display(value))


def fault_conclusion_text(value: Any, empty: str = "故障类型待复核") -> str:
    """将内部故障类型转换为面向用户的完整结论。"""
    text = display(value, empty).replace("预判", "判别")
    return text if text.endswith("故障") or text == empty else f"{text}故障"


def translate_wave_text(value: Any, empty: str = "未提供") -> str:
    """将行波模块返回的固定英文状态和原因转换为中文报告用语。"""
    if not has_value(value):
        return empty
    text = str(value).strip()
    exact = {
        "NotRequested": "未请求复核",
        "ModelOnly": "仅采用模型结果",
        "ConfirmedWithSupplement": "补充信息复核一致",
        "ReturnedInformationConfirmed": "回传信息复核一致",
        "ReturnedInformationCorrected": "已根据回传信息修正",
        "ExternalInformationConflict": "外部信息存在冲突",
        "ExternalConstraintCorrected": "已根据外部约束修正",
        "DistortedWaveformNeedsReview": "行波波形疑似畸变，需人工复核",
        "SubtypeUnresolved": "故障类型不能唯一确定",
        "SubtypeBoundaryNeedsReview": "普通绕击与雷击避雷线和导线型绕击边界待复核",
        "CandidateConflictNeedsReview": "候选结果存在冲突，需人工复核",
        "Above": "明显高于反击耐雷水平",
        "Below": "明显低于反击耐雷水平",
        "Gray": "接近耐雷水平",
        "GrayZone": "接近耐雷水平",
        "Missing": "未提供",
        "Consistent": "相别一致",
        "NotChecked": "未核对相别",
        "MainPhaseOutsideRecordedFaultPhase": "行波主相不在录波故障相范围内",
    }
    if text in exact:
        return exact[text]
    replacements = {
        "Recorded ABC fault restricts candidates to F/N": "录波故障相为ABC，允许候选限定为F/N",
        "Amplitude is clearly above withstand: candidates F/FBF": "雷电流幅值明显高于反击耐雷水平，允许候选为F/FBF",
        "Amplitude is clearly below withstand: candidates N/SF/FBF/FSF": "雷电流幅值明显低于反击耐雷水平，允许候选为N/SF/FBF/FSF",
        "The LN model returned N, but trusted recorded-phase three-wave evidence supports an allowed lightning subtype": "雷击/非雷击模型首次判别为N，但录波故障相对应的三相行波特征支持允许候选中的雷击绕击亚型",
        "Waveform hypothesis is inside the returned-information candidates": "行波特征判别结果位于回传信息允许的候选集合内",
        "No returned information; retained model result": "未提供回传信息，保留模型结果",
        "Phase and amplitude constraints have no common candidate": "相别与幅值约束没有共同候选",
        "Waveform is low-trust; unique external candidate was used": "行波可信度较低，采用外部信息限定的唯一候选",
        "Waveform is low-trust and external constraints are not unique": "行波可信度较低，且外部约束未形成唯一候选",
        "Known-phase three-wave evidence selected an allowed subtype": "已知故障相对应的三相行波特征选出了允许的绕击亚型",
        "R model and returned constraints leave one R subtype": "绕击模型与回传约束共同限定出唯一绕击亚型",
        "Generic three-phase subtype is inside allowed R candidates": "三相行波特征得到的亚型位于允许的绕击候选集合内",
        "R remains possible, but its subtype is not uniquely resolved": "仍可能属于绕击，但尚不能唯一确定绕击亚型",
        "Recorded ABC excludes R and the LN model supports lightning, so F was selected": "录波ABC三相故障排除绕击，且雷击/非雷击模型支持雷击，因此判定为反击",
        "Below-withstand constraint excludes F; three-phase evidence selected an R subtype": "低于反击耐雷水平的约束排除反击，三相行波特征选出绕击亚型",
        "Returned information leaves a unique candidate": "回传信息限定出唯一候选",
        "Model result is outside the candidates and review evidence is insufficient": "模型结果不在候选集合内，且复核证据不足",
        "Waveform lies on the SF/FBF boundary and tends toward FBF; subtype review is required": "行波特征处于普通绕击与雷击避雷线和导线型绕击的边界，且倾向后者，需人工复核故障类型",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(
        r"Amplitude ratio ([0-9.]+) is within the \+/-([0-9.]+)% gray zone; no type is removed",
        r"雷电流幅值接近耐雷水平（比值为\1，位于±\2%范围），不剔除候选类型",
        text,
    )
    text = re.sub(
        r"Traveling-wave main phase ([ABC]) is outside recorded phase ([ABC]+); waveform evidence is treated as distorted/low-trust",
        r"行波主相\1不在录波故障相\2范围内，行波证据按疑似畸变、低可信处理",
        text,
    )
    return text.replace("; ", "；").replace(";", "；")


def user_facing_model_terms(value: Any) -> str:
    """清理面向业务人员报告中的模型内部简称。"""
    text = str(value or "")
    phrase_replacements = {
        "SF_FBF_REVIEW": "普通绕击与雷击避雷线和导线型绕击边界待复核（倾向后者）",
        "SF/FBF待复核（倾向FBF）": "普通绕击与雷击避雷线和导线型绕击边界待复核（倾向后者）",
        "普通绕击（SF）": "普通绕击",
        "FSF型绕击（雷击多相导线）": "雷击多相导线型绕击",
        "FBF型绕击（雷击避雷线与导线）": "雷击避雷线与导线型绕击",
        "（R）": "",
        "（F）": "",
        "（N）": "",
        "（L）": "",
        "行波有效相": "行波候补故障相",
    }
    for source, target in phrase_replacements.items():
        text = text.replace(source, target)
    labels = {
        "FBF": "雷击避雷线与导线型绕击",
        "FSF": "雷击多相导线型绕击",
        "SF": "普通绕击",
        "R": "雷电绕击",
        "F": "雷电反击",
        "N": "非雷击",
        "L": "雷击",
    }
    for source, target in labels.items():
        text = re.sub(
            rf"(?<![A-Za-z0-9]){source}(?![A-Za-z0-9])",
            target,
            text,
        )
    return text


def usable_text(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return bool(text) and text not in {"default station", "default", "unknown", "未提供"}


def set_cell_fill(cell: Any, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell: Any, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_runs(paragraph: Any, size: float = 9.5, bold: bool = False, color: str | None = None) -> None:
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def normalize_document_fonts(doc: Document) -> None:
    """强制正文、表格、页眉页脚和样式统一使用微软雅黑。"""
    for style in doc.styles:
        try:
            style.font.name = "Microsoft YaHei"
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.get_or_add_rFonts()
            for key, value in (
                ("ascii", "Microsoft YaHei"),
                ("hAnsi", "Microsoft YaHei"),
                ("eastAsia", "微软雅黑"),
                ("cs", "Microsoft YaHei"),
            ):
                r_fonts.set(qn(f"w:{key}"), value)
        except (AttributeError, TypeError):
            continue
    for paragraph in iter_document_paragraphs(doc):
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
            r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
            r_fonts.set(qn("w:eastAsia"), "微软雅黑")
            r_fonts.set(qn("w:cs"), "Microsoft YaHei")


def iter_document_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[float]) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for index, text in enumerate(headers):
        header.cells[index].text = str(text)
        set_cell_fill(header.cells[index], PALE)
        style_runs(header.cells[index].paragraphs[0], bold=True, color=DARK)
    for values in rows:
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, value in enumerate(values):
            row.cells[index].text = display(value)
            style_runs(row.cells[index].paragraphs[0])
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _normalize_table_title(value: Any) -> str:
    """归一化表标题用于匹配：去除首部编号前缀与空白。"""
    title = str(value or "").strip()
    title = re.sub(r"^(表|附表)\s*([A-Za-z]?[0-9.]+)\s*", "", title)
    return re.sub(r"\s+", "", title)


def find_section_table(sections: Sequence[Mapping[str, Any]] | None, table_title: str) -> Mapping[str, Any] | None:
    """在 sections 富结构中查找标题匹配的 table 内容项；未找到返回 None。"""
    if not sections:
        return None
    target = _normalize_table_title(table_title)
    for section in sections:
        for item in section.get("content", []) or []:
            if isinstance(item, Mapping) and item.get("type") == "table":
                if _normalize_table_title(item.get("title")) == target:
                    return item
    return None


def table_title_from_spec(spec: Mapping[str, Any]) -> str:
    """从 table 内容项读取展示用标题（用于 docx 表题）。"""
    return str(spec.get("title") or "附表")


def _table_widths_by_columns(title: str, columns: Sequence[str]) -> list[float]:
    """按列数分配 docx 表格列宽（与既有报告版式保持一致）。"""
    col_count = len(columns)
    if col_count == 4:
        return [3.0, 5.5, 3.0, 5.5]
    if col_count == 3:
        return [3.5, 10.0, 3.5]
    if col_count == 2:
        return [4.2, 12.8]
    width = 17.0 / col_count
    return [width] * col_count


def add_table_from_spec(doc: Document, spec: Mapping[str, Any], caption: str | None = None) -> Any:
    """根据 sections 的 table 内容项渲染 docx 表格。"""
    columns: list[str] = [str(col) for col in (spec.get("columns") or [])]
    rows: list[list[Any]] = []
    for row in spec.get("rows") or []:
        cells = [str(cell) for cell in row]
        # 与列数对齐：不足补空、超出截断，防止 historic 数据绘制异常。
        if len(cells) < len(columns):
            cells = cells + [""] * (len(columns) - len(cells))
        elif len(cells) > len(columns):
            cells = cells[: len(columns)]
        rows.append(cells)
    if caption:
        add_table_caption(doc, caption)
    else:
        add_table_caption(doc, table_title_from_spec(spec))
    return add_table(doc, columns, rows, _table_widths_by_columns(table_title_from_spec(spec), columns))


def add_table_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(DARK)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_report_paragraph(doc: Document, text: str, *, first_line: bool = True, before: float = 0, after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.6
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.add_run(str(text))
    style_runs(paragraph, size=11)


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, 140, 160, 140, 160)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(f"{title}\n")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK)
    paragraph.add_run(text)
    style_runs(paragraph, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout_lines(doc: Document, title: str, lines: Sequence[str], fill: str = LIGHT) -> None:
    """生成分段式输出框，避免多项结论挤在同一段中。"""
    content = [str(line).strip() for line in lines if has_value(line)] or ["未提供"]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, 140, 160, 140, 160)
    first = cell.paragraphs[0]
    title_run = first.add_run(f"{title}\n")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(DARK)
    first.add_run(content[0])
    style_runs(first, size=10)
    for line in content[1:]:
        paragraph = cell.add_paragraph(line)
        style_runs(paragraph, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: Sequence[Any]) -> None:
    for item in items:
        if has_value(item):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(str(item).rstrip("。；") + "。")


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    paragraph.add_run(" 页")


def build_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color in (("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "输电线路故障综合判别系统｜正式报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_runs(header, size=8, color="666666")
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)
    style_runs(footer, size=8, color="777777")


def actual_sources(result: Mapping[str, Any]) -> list[str]:
    inputs = result.get("inputs", {})
    sources = ["录波故障数据"]
    if inputs.get("lightning"):
        sources.append("雷电定位数据")
    if safe_get(inputs, "line_ledger", "matched", default=False):
        sources.append("线路台账与风险评估数据")
    if inputs.get("traveling_wave_files"):
        sources.append("分布式行波数据")
    if inputs.get("optical"):
        sources.append("雷电光学监测数据")
    return sources


def resistance_sentence(result: Mapping[str, Any]) -> str:
    decision = result.get("ground_fault_decision", {})
    resistance = result.get("transition_resistance", {})
    if resistance.get("status") == "not_applicable" or decision.get("is_ground_fault") is False:
        return ""
    if not has_value(resistance.get("transition_resistance_ohm")):
        return "现有参数未能得到有效过渡电阻结果，建议复核电气量单位、相量关系及录波字段定义。"
    rf = float(resistance["transition_resistance_ohm"])
    if rf < 0:
        return (
            f"计算得到的过渡电阻为{format_number(rf)} Ω，结果小于0 Ω，属于物理异常结果；"
            "不得将其截断为0 Ω，也不得归入金属性、低阻或高阻接地，建议复核电气量单位、相量关系及零序电流定义。"
        )
    if resistance.get("status") != "ok":
        return "现有参数未能得到有效过渡电阻结果，建议复核电气量单位、相量关系及录波字段定义。"
    category = "金属性接地" if rf < 3 else "低阻接地" if rf < 500 else "高阻接地"
    criterion = "小于3 Ω" if rf < 3 else "处于3～500 Ω" if rf < 500 else "不小于500 Ω"
    return f"计算得到的过渡电阻为{format_number(rf)} Ω，{criterion}，属于{category}。"


def build_executive_summary(doc: Document, result: Mapping[str, Any]) -> None:
    inputs = result.get("inputs", {})
    fault = inputs.get("fault_record", {})
    lightning = inputs.get("lightning", {})
    ground = result.get("ground_fault_summary", {})
    recommendation = result.get("overall_recommendation", {})
    location = result.get("fault_location", {})
    ledger_record = safe_get(inputs, "line_ledger", "record", default={}) or {}
    lightning_cause = result.get("lightning_cause_analysis", {})
    wave = result.get("traveling_wave_analysis", {})
    fusion = result.get("cause_fusion", {})
    optical = inputs.get("optical", {})
    parts = [display(fault.get("fault_time"))]
    if usable_text(fault.get("station")):
        parts.append(str(fault["station"]))
    line_name = fault.get("line_name") or lightning.get("line_name")
    if has_value(line_name):
        parts.append(str(line_name))
    reclose = display(fault.get("reclosing"))
    first = " ".join(parts) + f"跳闸，重合闸{reclose}，录波显示为{fault_conclusion_text(ground.get('conclusion'))}。"
    if safe_get(result, "ground_fault_decision", "is_ground_fault") is True:
        first += resistance_sentence(result)
    second = "现结合" + "、".join(actual_sources(result)) + "对线路故障原因进行分析。"
    source_labels = "、".join(
        display(item.get("label"))
        for item in location.get("sources", [])
        if item.get("available")
    ) or "现有定位信息"
    location_sentence = (
        f"1）故障定位：{source_labels}综合显示故障位置为{display(location.get('result_text'))}，"
        f"对应位置地形为{display(ledger_record.get('terrain'))}。"
    )
    has_strike = str(lightning.get("query_result", "")).strip().upper() in {"是", "有", "YES", "Y", "TRUE", "1"}
    if has_strike:
        mismatch = any(
            item.get("code") == "event_time_mismatch"
            for item in safe_get(result, "input_validation", "issues", default=[])
            if isinstance(item, Mapping)
        )
        location_sentence += (
            f"雷电定位系统存在落雷记录，雷电流幅值为{value_with_unit(lightning.get('current_ka'), 'kA')}。"
            + (
                "雷电定位位置与测距结果存在空间关联，但事件时刻不一致，不能自动确认为同一故障事件。"
                if mismatch
                else "雷电定位位置与故障测距信息基本吻合。"
            )
        )
    else:
        location_sentence += "雷电定位系统未提供可用落雷记录。"
    cause_sentence = (
        f"2）故障成因：{display(lightning_cause.get('report_text'), '耐雷水平数据不足')}；"
        f"行波分析结果显示：{wave_fault_conclusion(result)}；"
        f"综合原因结论为：{fused_fault_conclusion(result)}。"
    )
    if optical:
        cause_sentence += f"光学监测结果为{display(optical.get('result') or optical.get('conclusion'))}。"
    reliability = f"可靠性与后续意见：{effective_review_summary(result)}。"
    follow_up = f"3）后续意见：{effective_review_decision(result)}。"
    add_heading(doc, "故障分析", 2)
    for text in (first, second, location_sentence, cause_sentence, reliability, follow_up):
        add_report_paragraph(doc, text)


def build_basic_information(doc: Document, result: Mapping[str, Any], front_only: bool, sections: Sequence[Mapping[str, Any]] | None = None) -> None:
    inputs = result.get("inputs", {})
    fault = inputs.get("fault_record", {})
    lightning = inputs.get("lightning", {})
    ground = result.get("ground_fault_summary", {})
    recommendation = result.get("overall_recommendation", {})
    resistance = result.get("transition_resistance", {})
    location = result.get("fault_location", {})
    cause = result.get("cause_fusion", {})
    lightning_cause = result.get("lightning_cause_analysis", {})
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    line_name = fault.get("line_name") or lightning.get("line_name")
    station_display = fault.get("station") if usable_text(fault.get("station")) else "未提供"
    resistance_category = (
        "不适用"
        if resistance.get("status") == "not_applicable"
        else display(resistance.get("category_report_text"), "无法计算")
    )
    overview = find_section_table(sections, "表1 故障分析总览")
    if overview is not None:
        add_table_from_spec(doc, overview, caption="表1 故障分析总览")
    else:
        add_table_caption(doc, "表1 故障分析总览")
        add_table(doc, ["项目", "内容", "项目", "内容"], [
            ["报告编号", report_identifier(result), "生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["故障线路", display(line_name), "故障时刻", display(fault.get("fault_time"))],
            ["录波站", station_display, "故障相", normalize_phase(fault.get("fault_phase")) or "未提供"],
            ["故障性质", display(ground.get("conclusion"), "待复核"), "接地类型", resistance_category],
            ["综合定位", display(location.get("result_text")), "雷击原因", fused_fault_conclusion(result)],
            ["行波结论", wave_fault_conclusion(result), "总体状态", effective_review_decision(result)],
        ], [2.4, 6.1, 2.4, 6.1])


def build_ground_and_resistance_section(doc: Document, result: Mapping[str, Any]) -> None:
    add_heading(doc, "2  接地故障及过渡电阻判别")
    summary = result.get("ground_fault_summary", {})
    resistance = result.get("transition_resistance", {})
    observations = "、".join(str(item).rstrip("。；") for item in summary.get("observations", []) if has_value(item))
    grounding_sentence = (
        f"录波特征显示{observations}，属于{fault_conclusion_text(summary.get('conclusion'))}。"
        if observations
        else f"根据录波信息，属于{fault_conclusion_text(summary.get('conclusion'))}。"
    )
    output_text = (
        grounding_sentence
        + resistance_sentence(result)
    )
    add_report_paragraph(doc, output_text)
    if resistance.get("status") == "not_applicable":
        return
    comparisons = resistance.get("zero_sequence_comparison", {})
    if comparisons:
        rows = []
        for candidate in comparisons.values():
            rows.append([
                candidate.get("assumption"),
                format_number(candidate.get("three_i0_a_used")),
                format_number(candidate.get("zm_ohm")),
                format_number(candidate.get("transition_resistance_ohm")),
                "有效" if candidate.get("status") == "ok" else "需复核",
                "采用" if candidate.get("selected") else "对照",
            ])
        add_table(doc, ["零序电流解释（仅单相接地）", "代入3I₀(A)", "Zm(Ω)", "Rf(Ω)", "有效性", "采用情况"], rows, [4.2, 2.6, 2.4, 2.6, 2.6, 2.6])
        add_report_paragraph(doc, "说明：两套结果用于核对零序电流字段语义，主结果依据录波字段定义选择，不能只根据结果正负反推零序电流含义。", first_line=False)


def build_location_section(doc: Document, result: Mapping[str, Any], sections: Sequence[Mapping[str, Any]] | None = None) -> None:
    add_heading(doc, "3  故障定位")
    location = result.get("fault_location", {})
    ledger = safe_get(result, "inputs", "line_ledger", default={}) or {}
    record = ledger.get("record", {}) if isinstance(ledger, Mapping) else {}
    available = [item for item in location.get("sources", []) if item.get("available")]
    rows = []
    for item in available:
        distance_text = (
            f"{format_number(item.get('distance_km'))} km"
            if has_value(item.get("distance_km"))
            else "位置记录可用"
        )
        rows.append([item.get("label"), distance_text, item.get("tower")])
    labels = "、".join(display(item.get("label")) for item in available) or "现有定位信息"
    terrain = display(record.get("terrain"))
    location_sentence = f"{labels}综合显示故障位置为{display(location.get('result_text'))}，对应位置地形为{terrain}。"
    lightning = safe_get(result, "inputs", "lightning", default={}) or {}
    has_strike = str(lightning.get("query_result", "")).strip().upper() in {"是", "有", "YES", "Y", "TRUE", "1"}
    if has_strike:
        validation_errors = [item.get("code") for item in safe_get(result, "input_validation", "issues", default=[]) if isinstance(item, Mapping)]
        if "event_time_mismatch" in validation_errors:
            lightning_sentence = f"查询雷电定位系统，线路跳闸时刻附近存在落雷记录，雷电流幅值为{format_number(lightning.get('current_ka'))} kA。雷电定位位置与测距结果存在空间关联，但事件时刻不一致，不能自动确认为同一故障事件，建议人工复核。"
        else:
            lightning_sentence = f"查询雷电定位系统，线路跳闸时刻附近存在落雷记录，雷电流幅值为{format_number(lightning.get('current_ka'))} kA，雷电定位位置与故障测距信息基本吻合。"
    else:
        lightning_sentence = "查询雷电定位系统，线路跳闸时刻附近无可用落雷记录。"
    add_report_paragraph(doc, location_sentence + lightning_sentence)
    sources_table = find_section_table(sections, "表2 定位信息来源")
    if sources_table is not None:
        add_table_from_spec(doc, sources_table, caption="表2 定位信息来源")
    elif rows:
        add_table_caption(doc, "表2 定位信息来源")
        add_table(doc, ["定位来源", "定位结果", "对应杆塔/区段"], rows, [3.5, 6.0, 7.5])


def build_lightning_section(doc: Document, result: Mapping[str, Any]) -> None:
    add_heading(doc, "4  耐雷水平与雷击原因分析")
    inputs = result.get("inputs", {})
    ledger = inputs.get("line_ledger", {})
    record = ledger.get("record", {}) if ledger.get("matched") else {}
    cause = result.get("lightning_cause_analysis", {})
    current = cause.get("lightning_current_ka") or safe_get(inputs, "lightning", "current_ka")
    if record:
        ledger_sentence = (
            f"经线路台账匹配，故障候选位置对应{display(record.get('tower_no'))}杆塔，"
            f"塔形为{display(record.get('tower_shape'))}，地形为{display(record.get('terrain'))}，"
            f"海拔{format_number(record.get('elevation_m'))} m，接地电阻实测值为"
            f"{format_number(record.get('ground_resistance_ohm'))} Ω。该塔反击耐雷水平为"
            f"{format_number(record.get('backflashover_withstand_ka'))} kA，绕击耐雷水平为"
            f"{format_number(record.get('shielding_failure_withstand_ka'))} kA。"
        )
    else:
        ledger_sentence = "本次未匹配到可用的故障杆塔台账，杆塔结构及耐雷水平参数暂无法核实。"
    current_sentence = (
        f"本次雷电流幅值为{format_number(current)} kA。"
        if has_value(current)
        else "本次未提供可用的雷电流幅值。"
    )
    conclusion_sentence = (
        f"根据雷电流幅值与台账耐雷水平的固定阈值比较，{display(cause.get('report_text'), '耐雷水平数据不足')}。"
    )
    add_report_paragraph(doc, ledger_sentence + current_sentence + conclusion_sentence)


def subtype_text(value: Any) -> str:
    return {
        "SF": "普通绕击",
        "FSF": "雷击多相导线型绕击",
        "FBF": "雷击避雷线与导线型绕击",
        "SF_FBF_REVIEW": "普通绕击与雷击避雷线和导线型绕击边界待复核（倾向后者）",
    }.get(str(value or "").upper(), "故障类型不能唯一确定")


def wave_candidate_types_display(value: Any) -> str:
    """将内部候选标签转换为面向报告阅读者的中文名称。"""
    if not has_value(value):
        return "未形成可用候选集合"
    mapping = {
        "N": "非雷击",
        "F": "雷电反击",
        "R": "雷电绕击",
        "SF": "普通绕击",
        "FBF": "雷击避雷线与导线型绕击",
        "FSF": "雷击多相导线型绕击",
        "SF_FBF_REVIEW": "普通绕击与雷击避雷线和导线型绕击边界待复核（倾向后者）",
    }
    tokens = [item.strip().upper() for item in re.split(r"[/,，、;；\s]+", str(value)) if item.strip()]
    return "、".join(mapping.get(item, item) for item in tokens) or "未形成可用候选集合"


def wave_review_trigger_text(model: Mapping[str, Any], review: Mapping[str, Any]) -> str:
    if model.get("correction_applied"):
        return "雷电幅值判别结果、录波故障相与行波特征有出入，请求复核"
    if model.get("review_confirmed"):
        return "雷电幅值判别结果、录波故障相与行波特征一致，完成复核"
    if review.get("review_triggered"):
        return "雷电幅值判别结果、录波故障相与行波特征存在待核项，请求复核"
    return "雷电幅值判别结果、录波故障相与行波特征未触发额外复核"


def returned_information_sentence(
    wave: Mapping[str, Any], model: Mapping[str, Any], candidate_text: str
) -> str:
    phase = sequence_phase(model.get("corrected_fault_phase") or wave.get("recording_fault_phase"))
    prefix = f"程序回传录波故障相{phase}及幅值信息复核" if phase else "程序回传幅值信息复核"
    band = str(model.get("amplitude_band") or "")
    if band in {"Above", "Below", "Gray", "GrayZone"} and has_value(candidate_text):
        return f"{prefix}，因雷电流幅值{translate_wave_text(band)}，允许候选为{candidate_text}。"
    return f"{prefix}。"


def wave_ln_basis_sentence(wave: Mapping[str, Any], model: Mapping[str, Any]) -> str:
    label = str(
        model.get("ln_prediction")
        or model.get("initial_final_label")
        or model.get("final_label")
        or wave.get("initial_classification")
        or ""
    ).upper()
    classification = str(wave.get("initial_classification") or "")
    lightning_score = model.get("ln_lightning_score")
    non_lightning_score = model.get("ln_non_lightning_score")
    if label == "N" or "非雷击" in classification:
        return "行波波尾拖尾较长、低频成分占比较高，整体特征偏向非雷击故障。"
    if label in {"L", "F", "R", "SF", "FSF", "FBF"} or "雷击" in classification:
        return "行波首波波前较陡、暂态高频成分较明显，整体特征偏向雷击故障。"
    if has_value(lightning_score) and has_value(non_lightning_score):
        try:
            if float(lightning_score) >= float(non_lightning_score):
                return "行波首波波前较陡、暂态高频成分较明显，整体特征偏向雷击故障。"
            return "行波波尾拖尾较长、低频成分占比较高，整体特征偏向非雷击故障。"
        except (TypeError, ValueError):
            pass
    return "行波雷击与非雷击特征区分度不足，需结合录波故障相和雷电定位信息复核。"


def normalized_two_class_probabilities(first: Any, second: Any) -> tuple[float, float] | None:
    """将两类模型判别值经稳定Softmax归一化为工程展示概率。"""
    if not has_value(first) or not has_value(second):
        return None
    try:
        first_value = float(first)
        second_value = float(second)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(first_value) or not math.isfinite(second_value):
        return None
    maximum = max(first_value, second_value)
    first_exp = math.exp(first_value - maximum)
    second_exp = math.exp(second_value - maximum)
    total = first_exp + second_exp
    return (first_exp / total, second_exp / total) if total else None


def wave_fr_probability_sentence(model: Mapping[str, Any]) -> str:
    """将反击/绕击原始判别值转换为便于工程阅读的归一化概率。"""
    if str(model.get("initial_final_label") or "").upper() == "N":
        return "未进入反击与绕击判别环节"
    stored_backflash = model.get("fr_backflash_probability")
    stored_shielding = model.get("fr_shielding_probability")
    probabilities = (
        (float(stored_backflash), float(stored_shielding))
        if has_value(stored_backflash) and has_value(stored_shielding)
        else normalized_two_class_probabilities(
            model.get("fr_backflash_score"), model.get("fr_shielding_score")
        )
    )
    if probabilities:
        backflash_probability, shielding_probability = probabilities
        return (
            f"模型判断反击概率{backflash_probability * 100:.2f}%，"
            f"绕击概率{shielding_probability * 100:.2f}%"
        )
    return "反击与绕击概率未提供"


def wave_winning_ln_probability_sentence(wave: Mapping[str, Any], model: Mapping[str, Any]) -> str:
    label = str(model.get("ln_prediction") or model.get("initial_final_label") or "").upper()
    classification = str(wave.get("initial_classification") or "")
    if label == "N" or "非雷击" in classification:
        score = model.get("ln_non_lightning_score")
        return f"模型判断非雷击概率{format_number(float(score) * 100, 2)}%" if has_value(score) else "模型未提供非雷击概率"
    score = model.get("ln_lightning_score")
    return f"模型判断雷击概率{format_number(float(score) * 100, 2)}%" if has_value(score) else "模型未提供雷击概率"


def wave_winning_fr_probability_sentence(wave: Mapping[str, Any], model: Mapping[str, Any]) -> str:
    label = str(model.get("fr_prediction") or model.get("initial_final_label") or "").upper()
    if label == "N" or "非雷击" in str(wave.get("initial_classification") or ""):
        return ""
    stored_backflash = model.get("fr_backflash_probability")
    stored_shielding = model.get("fr_shielding_probability")
    probabilities = (
        (float(stored_backflash), float(stored_shielding))
        if has_value(stored_backflash) and has_value(stored_shielding)
        else normalized_two_class_probabilities(
            model.get("fr_backflash_score"), model.get("fr_shielding_score")
        )
    )
    if not probabilities:
        return "模型未提供反击与绕击概率"
    backflash_probability, shielding_probability = probabilities
    if label == "F" or "反击" in str(wave.get("initial_classification") or ""):
        return f"模型判断反击概率{backflash_probability * 100:.2f}%"
    return f"模型判断绕击概率{shielding_probability * 100:.2f}%"


def wave_fr_basis_sentence(wave: Mapping[str, Any], model: Mapping[str, Any]) -> str:
    """根据首次F/R标签及模型判别值生成概括性的反击/绕击特征依据。"""
    label = str(
        model.get("fr_prediction")
        or model.get("initial_final_label")
        or wave.get("initial_classification")
        or ""
    ).upper()
    classification = str(wave.get("initial_classification") or "")
    if label == "N" or "非雷击" in classification:
        return "本次在雷击与非雷击环节判为非雷击，未进入反击与绕击判别环节。"
    if label == "F" or "反击" in classification:
        return (
            "同时，局部波形起伏和过零活动较多、主峰频谱重心及高频能量较高、"
            "峰前极性脉冲活动较明显，整体特征偏向雷电反击故障。"
        )
    if label in {"R", "SF", "FBF", "FSF"} or "绕击" in classification:
        return (
            "同时，局部起伏和过零活动较少、主峰频谱的中频能量较高、"
            "峰前脉冲活动较弱，整体特征偏向雷电绕击故障。"
        )
    return "反击与绕击特征区分度不足，需结合录波故障相、雷电流幅值及原始行波复核。"


def wave_review_result_sentence(
    wave: Mapping[str, Any], model: Mapping[str, Any], review: Mapping[str, Any]
) -> str:
    final_phase = sequence_phase(model.get("corrected_fault_phase") or wave.get("recording_fault_phase"))
    wave_phase = wave_analysis_phase(wave)
    final_label = str(model.get("final_label") or "").upper()
    subtype = str(model.get("final_subtype") or "").upper()
    correction_status = str(model.get("correction_status") or "")
    recording_phase = sequence_phase(review.get("recording_fault_phase") or final_phase)
    active_phases = review.get("initial_wave_active_phases") or model.get("active_phases")
    if isinstance(active_phases, Sequence) and not isinstance(active_phases, str):
        active_phases = "".join(str(item) for item in active_phases)
    active_phases = sequence_phase(active_phases)
    prefix = "结合录波故障相及雷电流幅值判别结果进行复核，复核结果显示"

    if correction_status == "SubtypeUnresolved":
        return (
            f"{prefix}现有行波特征未能满足唯一故障类型判据，故障类型不能唯一确定，"
            "建议人工复核原始三相行波。"
        )
    if correction_status == "SubtypeBoundaryNeedsReview" or subtype == "SF_FBF_REVIEW":
        return (
            f"{prefix}{final_phase or recording_phase or '相关'}相行波第二强相与主相幅值比处于普通绕击和"
            "雷击避雷线与导线型绕击的边界区间，且主峰后存在较明显的多峰活动，"
            "故障类型倾向雷击避雷线与导线型绕击，但尚不能直接确定，建议人工复核。"
        )
    if correction_status == "DistortedWaveformNeedsReview":
        return (
            f"{prefix}行波判别故障相{wave_phase or '未提供'}与录波故障相{recording_phase or '未提供'}不一致，"
            "行波相别证据按疑似波形畸变处理，现有特征不足以可靠修正故障类型，建议人工复核。"
        )
    if correction_status in {"CandidateConflictNeedsReview", "ExternalInformationConflict"}:
        return (
            f"{prefix}雷电流幅值对应的候选类型、录波故障相约束与行波模型结果之间存在冲突，"
            "未形成共同且唯一的故障类型，建议人工复核。"
        )
    if review.get("phase_matches") is False:
        return (
            f"{prefix}行波候补故障相{active_phases or '未提供'}与录波故障相{recording_phase or '未提供'}不一致，"
            "回传录波故障相复核后仍未满足相别一致条件，建议人工复核原始行波。"
        )
    if review.get("amplitude_matches") is False:
        return (
            f"{prefix}雷电流幅值与耐雷水平比较得到的故障类型与行波首次判别类型不一致，"
            "回传幅值信息复核后仍未形成唯一结论，建议人工复核。"
        )
    if final_label == "N":
        return f"{prefix}行波波尾拖尾较长、低频成分占比较高，整体特征偏向非雷击故障。"
    if subtype == "FSF" and len(final_phase) >= 2:
        phase_count = "两相" if len(final_phase) == 2 else "三相"
        return (
            f"{prefix}{final_phase}{phase_count}行波偏向两相同时明显，两个主要相幅值接近、"
            "主峰极性相同、高频能量比例接近，符合雷击多相导线型绕击判据。"
        )
    if subtype == "FBF" and final_phase:
        return (
            f"{prefix}{final_phase}相偏向单相主导，并且主峰后有长波尾、多峰活动或较强后续能量，"
            "符合雷击避雷线与导线型绕击判据。"
        )
    if subtype == "SF" and final_phase:
        return (
            f"{prefix}{final_phase}相局部起伏和过零活动较少、主峰频谱的中频能量较高、"
            "峰前脉冲活动较弱，符合普通绕击判据。"
        )
    if final_label == "F" and wave_phase and final_phase:
        return (
            f"{prefix}行波判别故障相{wave_phase}位于录波故障相{final_phase}范围内，"
            "局部波形起伏和过零活动较多、主峰频谱重心及高频能量较高、"
            "峰前极性脉冲活动较明显，整体特征偏向雷电反击故障。"
        )
    if model.get("review_confirmed"):
        return f"{prefix}录波故障相、雷电流幅值判别结果与行波判别结果一致，保留初判结果。"
    if model.get("correction_applied"):
        reason = translate_wave_text(model.get("correction_reason"), "外部信息与首次行波结果不一致")
        return f"{prefix}由于{reason}，已根据录波故障相、雷电流幅值和行波特征完成修正。"
    if review.get("review_triggered"):
        reasons = review.get("trigger_reasons") or []
        if isinstance(reasons, str):
            reasons = [reasons]
        reason_text = "；".join(str(item).strip("。；") for item in reasons if has_value(item))
        if reason_text:
            return f"{prefix}存在以下不一致：{reason_text}；回传信息后仍未形成唯一故障类型，建议人工复核。"
        return f"{prefix}录波故障相、雷电流幅值约束与行波特征未形成共同且唯一的故障类型，建议人工复核。"
    return f"{prefix}录波故障相及雷电流幅值信息未触发行波结果修正。"


def build_wave_section(doc: Document, result: Mapping[str, Any], sections: Sequence[Mapping[str, Any]] | None = None) -> None:
    add_heading(doc, "5.1  分布式行波分析", 2)
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    active = review.get("initial_wave_active_phases") or model.get("active_phases")
    if isinstance(active, Sequence) and not isinstance(active, str):
        active = "/".join(str(item) for item in active)
    initial_phase = model.get("main_phase") or review.get("initial_wave_main_phase")
    phase_display = f"行波判别故障相{initial_phase}、候补故障相{active}" if has_value(initial_phase) or has_value(active) else "行波模型未输出可用的判别故障相或候补故障相信息"
    lightning_score = model.get("ln_lightning_score")
    non_lightning_score = model.get("ln_non_lightning_score")
    score_text = f"模型判断雷击概率{format_number(float(lightning_score) * 100, 2)}%，非雷击概率{format_number(float(non_lightning_score) * 100, 2)}%" if has_value(lightning_score) and has_value(non_lightning_score) else "雷击与非雷击概率未提供"
    fr_probability_text = wave_fr_probability_sentence(model)
    final_phase = model.get("corrected_fault_phase") or wave.get("recording_fault_phase")
    candidate_text = wave_candidate_types_display(model.get("candidate_types"))
    returned_sentence = returned_information_sentence(wave, model, candidate_text)
    ln_basis = wave_ln_basis_sentence(wave, model)
    fr_basis = wave_fr_basis_sentence(wave, model)
    initial_basis = f"{ln_basis}{fr_basis}"
    ln_winning_probability = wave_winning_ln_probability_sentence(wave, model)
    fr_winning_probability = wave_winning_fr_probability_sentence(wave, model)
    review_result = wave_review_result_sentence(wave, model, review)
    low_score_note = ""
    available_scores = [float(v) for v in (lightning_score, non_lightning_score) if has_value(v)]
    if available_scores and max(available_scores) < 0.60:
        low_score_note = "首次雷击与非雷击模型最高概率低于60%，模型区分度较低，建议结合原始行波复核。"
    add_report_paragraph(doc, wave_process_summary_sentence(result))
    wave_table = find_section_table(sections, "表3 行波判别过程")
    if wave_table is not None:
        add_table_from_spec(doc, wave_table, caption="表3 行波判别过程")
    else:
        add_table_caption(doc, "表3 行波判别过程")
        add_table(doc, ["项目", "结果"], [
            ["初判结果", f"{display(wave.get('initial_classification'))}；{phase_display}；{score_text}；{fr_probability_text}"],
            ["初判依据", initial_basis],
            ["允许候选", candidate_text],
            ["修正/复核结果", review_result],
            ["最终采用结果", wave_fault_conclusion(result)],
        ], [4.2, 12.8])
    note = "说明：反击与绕击概率为模型输出的工程归一化结果，用于辅助判别，不作为统计置信度使用。"
    if low_score_note:
        note += low_score_note
    add_report_paragraph(doc, note, first_line=False, before=2)


def recording_phase_sentence(result: Mapping[str, Any]) -> str:
    phase = corrected_fault_phase(result)
    return f"{phase}相故障" if phase else "故障相未提供"


def lightning_location_conclusion(result: Mapping[str, Any]) -> str:
    cause = result.get("lightning_cause_analysis", {})
    text = display(cause.get("report_text") or cause.get("conclusion") or cause.get("mechanism"), "雷电流幅值未提供")
    text = user_facing_model_terms(text).rstrip("。；")
    return text


def wave_initial_type_detail(result: Mapping[str, Any]) -> str:
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    label = str(model.get("final_label") or wave.get("classification") or "").upper()
    subtype = str(model.get("final_subtype") or "").upper()
    if label == "N" or "非雷击" in str(wave.get("initial_classification") or ""):
        return "非雷击"
    if label == "F" or "反击" in str(wave.get("initial_classification") or ""):
        return "雷电反击"
    if label == "R" or "绕击" in str(wave.get("initial_classification") or ""):
        return f"雷电绕击/{subtype_text(subtype)}" if subtype in {"SF", "FBF", "FSF", "SF_FBF_REVIEW"} else "雷电绕击"
    return "待复核"


def wave_review_summary_sentence(result: Mapping[str, Any]) -> str:
    return wave_initial_review_sentence(result)


def concise_wave_fault_conclusion(result: Mapping[str, Any]) -> str:
    conclusion = wave_fault_conclusion(result)
    return "可能是" + conclusion.split("，可能是", 1)[1] if "，可能是" in conclusion else conclusion


def wave_process_summary_sentence(result: Mapping[str, Any]) -> str:
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    parts = [
        wave_ln_basis_sentence(wave, model).rstrip("。；"),
        wave_fr_basis_sentence(wave, model).rstrip("。；"),
    ]
    probability_parts = [
        wave_winning_ln_probability_sentence(wave, model).rstrip("。；"),
        wave_winning_fr_probability_sentence(wave, model).rstrip("。；"),
    ]
    review_result = wave_review_result_sentence(wave, model, review).rstrip("。；")
    final_result = concise_wave_fault_conclusion(result).rstrip("。；")
    return "。".join(item for item in parts if has_value(item)) + "。" + "，".join(
        item for item in probability_parts + [review_result, final_result] if has_value(item)
    ) + "。"


def wave_initial_review_sentence(result: Mapping[str, Any]) -> str:
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    recording_phase = corrected_fault_phase(result)
    wave_phase = wave_conclusion_phase(result)
    candidate_phase = wave_candidate_phase(wave)
    subtype = str(model.get("final_subtype") or "").upper()
    correction_status = str(model.get("correction_status") or "")
    if correction_status == "SubtypeBoundaryNeedsReview" or subtype == "SF_FBF_REVIEW":
        phase = recording_phase or wave_phase or candidate_phase
        conductor = phase_conductor_text(phase)
        return (
            f"结合录波故障信息及雷电定位进行复核，本次故障可能为雷电绕击{conductor}故障；"
            "现有行波特征处于普通绕击与雷击避雷线和导线型绕击的边界，特征倾向后者，"
            "但尚不能直接确定故障类型，建议人工复核。"
        )
    if ordinary_shielding_multi_phase_conflict(result):
        phase_count = "两相" if len(recording_phase) == 2 else "三相"
        return (
            f"结合录波故障信息及雷电定位，录波检测为{recording_phase}{phase_count}故障，但行波仅符合普通绕击特征；"
            f"按多相故障融合规则，本次故障可判为{fused_fault_conclusion(result)}。"
            "由于现有行波特征未满足雷击多相导线型绕击判据，建议人工复核。"
        )
    has_phase_containment = phases_have_containment(recording_phase, wave_phase) or phases_have_containment(
        recording_phase, candidate_phase
    )
    if model.get("review_confirmed"):
        status = "通过"
    elif model.get("correction_applied"):
        status = "经复核修正后通过"
    elif review.get("review_triggered"):
        status = "未通过"
    else:
        status = "未触发额外复核"
    if has_phase_containment or model.get("review_confirmed") or model.get("correction_applied"):
        return f"结合录波故障信息及雷电定位的行波判别{status}，本次故障可判为{fused_fault_conclusion(result)}。"

    conflicts: list[str] = []
    if recording_phase and (wave_phase or candidate_phase):
        wave_phase_text = candidate_phase or wave_phase
        conflicts.append(f"录波故障相为{recording_phase}，行波判别故障相/候补故障相为{wave_phase_text}，二者不存在包含关系")
    if review.get("phase_matches") is False:
        conflicts.append("录波故障相与行波相别复核不一致")
    if review.get("amplitude_matches") is False:
        conflicts.append("雷电流幅值判别结果与行波故障类型不一致")
    correction_status = translate_wave_text(model.get("correction_status"), "")
    if correction_status and correction_status not in {"补充信息复核一致", "已根据回传信息修正", "仅采用模型结果"}:
        conflicts.append(correction_status)
    if not conflicts and review.get("trigger_reasons"):
        reasons = review.get("trigger_reasons")
        if isinstance(reasons, str):
            reasons = [reasons]
        conflicts.extend(translate_wave_text(item, "") for item in reasons if has_value(item))
    conflict_text = "；".join(dict.fromkeys(item for item in conflicts if has_value(item))) or "录波故障信息、雷电流幅值判别结果与行波特征未形成一致结论"
    return (
        f"结合录波故障信息及雷电定位的行波判别{status}，行波模型结论为{wave_fault_conclusion(result)}；"
        f"但{conflict_text}，建议人工复核。"
    )


def build_fusion_section(doc: Document, result: Mapping[str, Any], sections: Sequence[Mapping[str, Any]] | None = None) -> None:
    add_heading(doc, "5.2  故障原因融合", 2)
    wave = result.get("traveling_wave_analysis", {})
    rows = [
        ["录波故障相", recording_phase_sentence(result)],
        ["雷电定位", lightning_location_conclusion(result)],
        ["分布式行波", f"行波初判为{display(wave.get('initial_classification'))}（{wave_initial_type_detail(result)}）"],
    ]
    add_report_paragraph(doc, wave_initial_review_sentence(result))
    fusion_table = find_section_table(sections, "表4 故障原因融合分析")
    if fusion_table is not None:
        add_table_from_spec(doc, fusion_table, caption="表4 故障原因融合分析")
    else:
        add_table_caption(doc, "表4 故障原因融合分析")
        table = add_table(doc, ["信息来源", "结论"], rows, [4.2, 12.8])
        row = table.add_row()
        merged = row.cells[0].merge(row.cells[1])
        merged.text = wave_review_summary_sentence(result)
        style_runs(merged.paragraphs[0], size=9.5)
    add_report_paragraph(doc, "证据优先级为：光学监测明确证据、有效行波结论、雷电流幅值与台账耐雷水平比较。该优先级用于处理冲突，不用于生成概率。", first_line=False, before=2)


def build_recommendation_section(doc: Document, result: Mapping[str, Any]) -> None:
    add_heading(doc, "6  总体状态与后续意见")
    recommendation = result.get("overall_recommendation", {})
    add_report_paragraph(
        doc,
        f"{user_facing_model_terms(effective_review_decision(result))}："
        f"{user_facing_model_terms(effective_review_summary(result))}。",
    )
    add_bullets(doc, [user_facing_model_terms(item) for item in recommendation.get("reasons", [])])


def build_input_appendix(doc: Document, result: Mapping[str, Any], sections: Sequence[Mapping[str, Any]] | None = None) -> None:
    doc.add_page_break()
    add_heading(doc, "附表A  输入信息汇总")
    inputs = result.get("inputs", {})
    paths = result.get("source_paths", {})
    fault = inputs.get("fault_record", {})
    lightning = inputs.get("lightning", {})
    record_distance = fault.get("fault_distance_km") or lightning.get("recording_distance_km")
    zero_current = fault.get("zero_sequence_current")
    add_heading(doc, "A.1  录波故障信息", 2)
    t5 = find_section_table(sections, "表A.1 录波故障信息")
    if t5 is not None:
        add_table_from_spec(doc, t5, caption="表A.1 录波故障信息")
    else:
        add_table(doc, ["字段", "结果", "字段", "结果"], [
            ["故障时刻", fault.get("fault_time"), "故障线路", fault.get("line_name") or lightning.get("line_name")],
            ["故障相", normalize_fault_phase_code(fault.get("fault_phase")) or "未提供", "故障持续时间", value_with_unit(fault.get("duration_ms"), "ms")],
            ["重合闸情况", fault.get("reclosing"), "故障相电压", fault.get("phase_voltage")],
            ["故障相电流", fault.get("phase_current"), "提取零序电流", f"{display(zero_current)}（按3I₀/I₀双口径核对）"],
            ["最大零序电流", zero_current, "零序补偿系数", fault.get("zero_sequence_compensation_k")],
            ["故障测距", value_with_unit(record_distance, "km"), "原始文件", paths.get("fault")],
        ], [3.0, 5.5, 3.0, 5.5])
    add_heading(doc, "A.2  雷电流及定位信息", 2)
    t6 = find_section_table(sections, "表A.2 雷电流及定位信息")
    if t6 is not None:
        add_table_from_spec(doc, t6, caption="表A.2 雷电流及定位信息")
    else:
        add_table(doc, ["字段", "结果", "字段", "结果"], [
            ["雷电流幅值", value_with_unit(lightning.get("current_ka"), "kA"), "定位时间", lightning.get("location_time")],
            ["最近杆塔", lightning.get("nearest_tower"), "录波测距", value_with_unit(lightning.get("recording_distance_km"), "km")],
            ["行波测距", value_with_unit(lightning.get("traveling_wave_distance_km"), "km"), "定位距离", value_with_unit(lightning.get("distance_to_line_km"), "km")],
            ["输入文件", paths.get("lightning"), "数据状态", "已提供" if lightning else "未提供"],
        ], [3.0, 5.5, 3.0, 5.5])
    ledger = inputs.get("line_ledger", {})
    record = ledger.get("record", {}) if ledger.get("matched") else {}
    add_heading(doc, "A.3  线路台账匹配信息", 2)
    t7 = find_section_table(sections, "表A.3 线路台账匹配信息")
    if t7 is not None:
        add_table_from_spec(doc, t7, caption="表A.3 线路台账匹配信息")
    else:
        add_table(doc, ["字段", "结果", "字段", "结果"], [
            ["匹配线路/杆塔", f"{display(record.get('line_name'))} / {display(record.get('tower_no'))}", "塔形", record.get("tower_shape")],
            ["地形", record.get("terrain"), "海拔", value_with_unit(record.get("elevation_m"), "m")],
            ["绝缘子串长", value_with_unit(record.get("insulator_length_m"), "m"), "接地电阻", value_with_unit(record.get("ground_resistance_ohm"), "Ω")],
            ["绕击耐雷水平", value_with_unit(record.get("shielding_failure_withstand_ka"), "kA"), "反击耐雷水平", value_with_unit(record.get("backflashover_withstand_ka"), "kA")],
            ["风险评估等级", record.get("lightning_risk_level"), "台账文件", paths.get("ledger")],
            ["匹配状态", "已匹配" if ledger.get("matched") else "未匹配", "匹配说明", ledger.get("diagnostic") or ledger.get("reason") or "线路及杆塔匹配成功"],
        ], [3.0, 5.5, 3.0, 5.5])
    add_heading(doc, "A.4  行波与光学监测信息", 2)
    t8 = find_section_table(sections, "表A.4 行波与光学监测信息")
    if t8 is not None:
        add_table_from_spec(doc, t8, caption="表A.4 行波与光学监测信息")
    else:
        wave = result.get("traveling_wave_analysis", {})
        model = wave.get("model_result", {})
        wave_files = inputs.get("traveling_wave_files", [])
        wave_formats = "/".join(sorted({Path(str(item.get("file_name") or "")).suffix.lstrip(".").upper() for item in wave_files if item.get("file_name")})) or "未提供"
        optical = inputs.get("optical", {})
        initial_phase = model.get("main_phase") or safe_get(wave, "fault_phase_review", "initial_wave_main_phase")
        wave_input = (
            f"初判：{display(wave.get('initial_classification'))}；初判相：{display(initial_phase)}；"
            f"幅值信息：{translate_wave_text(safe_get(wave, 'fault_phase_review', 'amplitude_band'))}；格式：{wave_formats}"
        )
        optical_input = (
            f"位置：{display(optical.get('location') or optical.get('device_tower'))}；"
            f"现象/结论：{display(optical.get('result') or optical.get('conclusion'))}"
            if optical
            else "本次未提供光学监测数据"
        )
        add_table(doc, ["来源", "主要输入", "数据状态"], [
            ["分布式行波", wave_input, "已提供" if wave_files else "未提供"],
            ["光学监测", optical_input, "已提供" if optical else "未提供"],
        ], [3.0, 10.5, 3.5])


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def sanitize_report_model_terms(doc: Document) -> None:
    """对最终真实报告做兜底清理，避免内部模型简称进入用户文档。"""
    parts = [doc.part]
    for section in doc.sections:
        parts.extend((section.header.part, section.footer.part))
    seen: set[int] = set()
    for part in parts:
        marker = id(part)
        if marker in seen:
            continue
        seen.add(marker)
        for text_node in part.element.iter(qn("w:t")):
            if text_node.text:
                text_node.text = user_facing_model_terms(text_node.text)


def report_field_values(result: Mapping[str, Any]) -> dict[str, str]:
    """把判别结果转换为可填充模板字段；模板负责版式，脚本只负责字段值。"""
    inputs = result.get("inputs", {})
    paths = result.get("source_paths", {})
    fault = inputs.get("fault_record", {})
    lightning = inputs.get("lightning", {})
    ledger = inputs.get("line_ledger", {})
    ledger_record = ledger.get("record", {}) if ledger.get("matched") else {}
    optical = inputs.get("optical", {})
    ground = result.get("ground_fault_summary", {})
    resistance = result.get("transition_resistance", {})
    location = result.get("fault_location", {})
    lightning_cause = result.get("lightning_cause_analysis", {})
    wave = result.get("traveling_wave_analysis", {})
    model = wave.get("model_result", {})
    review = wave.get("fault_phase_review", {})
    recommendation = result.get("overall_recommendation", {})
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line_name = fault.get("line_name") or lightning.get("line_name")
    phase = normalize_phase(fault.get("fault_phase")) or "未提供"
    resistance_category = "不适用" if resistance.get("status") == "not_applicable" else display(
        resistance.get("category_report_text"), "无法计算"
    )

    event_parts = [display(fault.get("fault_time"))]
    if usable_text(fault.get("station")):
        event_parts.append(str(fault["station"]))
    if has_value(line_name):
        event_parts.append(str(line_name))
    event_sentence = " ".join(event_parts) + (
        f"跳闸，重合闸{display(fault.get('reclosing'))}，录波显示为"
        f"{fault_conclusion_text(ground.get('conclusion'))}。"
    )
    if safe_get(result, "ground_fault_decision", "is_ground_fault") is True:
        event_sentence += resistance_sentence(result)
    sources_sentence = "现结合" + "、".join(actual_sources(result)) + "对线路故障原因进行分析。"

    available_sources = [item for item in location.get("sources", []) if item.get("available")]
    source_labels = "、".join(display(item.get("label")) for item in available_sources) or "现有定位信息"
    terrain = display(ledger_record.get("terrain"))
    summary_location = (
        f"1）故障定位：{source_labels}综合显示故障位置为{display(location.get('result_text'))}，"
        f"对应位置地形为{terrain}。"
    )
    location_section = (
        f"{source_labels}综合显示故障位置为{display(location.get('result_text'))}，对应位置地形为{terrain}。"
    )
    has_strike = str(lightning.get("query_result", "")).strip().upper() in {"是", "有", "YES", "Y", "TRUE", "1"}
    mismatch = any(
        item.get("code") == "event_time_mismatch"
        for item in safe_get(result, "input_validation", "issues", default=[])
        if isinstance(item, Mapping)
    )
    if has_strike:
        relation = (
            "雷电定位位置与测距结果存在空间关联，但事件时刻不一致，不能自动确认为同一故障事件，建议人工复核。"
            if mismatch else "雷电定位位置与故障测距信息基本吻合。"
        )
        summary_location += f"雷电定位系统存在落雷记录，雷电流幅值为{value_with_unit(lightning.get('current_ka'), 'kA')}。{relation}"
        location_section += (
            f"查询雷电定位系统，线路跳闸时刻附近存在落雷记录，雷电流幅值为"
            f"{value_with_unit(lightning.get('current_ka'), 'kA')}，{relation}"
        )
    else:
        summary_location += "雷电定位系统未提供可用落雷记录。"
        location_section += "查询雷电定位系统，线路跳闸时刻附近无可用落雷记录。"

    summary_cause = (
        f"2）故障成因：{display(lightning_cause.get('report_text'), '耐雷水平数据不足')}；"
        f"行波分析结果显示：{wave_fault_conclusion(result)}；"
        f"综合原因结论为：{fused_fault_conclusion(result)}。"
    )
    if optical:
        summary_cause += f"光学监测结果为{display(optical.get('result') or optical.get('conclusion'))}。"

    observations = "、".join(str(item).rstrip("。；") for item in ground.get("observations", []) if has_value(item))
    observation_items = [str(item).rstrip("。；") for item in ground.get("observations", []) if has_value(item)]
    current_observation = observation_items[0] if len(observation_items) > 0 else f"{phase}相电流突增"
    voltage_observation = observation_items[1] if len(observation_items) > 1 else f"{phase}相电压明显跌落"
    zero_observation = observation_items[2] if len(observation_items) > 2 else "零序电流及零序电压需进一步复核"
    grounding_sentence = (
        f"录波特征显示{observations}，属于{fault_conclusion_text(ground.get('conclusion'))}。"
        if observations else f"根据录波信息，属于{fault_conclusion_text(ground.get('conclusion'))}。"
    ) + resistance_sentence(result)

    current = lightning_cause.get("lightning_current_ka") or lightning.get("current_ka")
    if ledger_record:
        lightning_section = (
            f"经线路台账匹配，故障候选位置对应{display(ledger_record.get('tower_no'))}杆塔，"
            f"塔形为{display(ledger_record.get('tower_shape'))}，地形为{display(ledger_record.get('terrain'))}，"
            f"海拔{format_number(ledger_record.get('elevation_m'))} m，接地电阻实测值为"
            f"{format_number(ledger_record.get('ground_resistance_ohm'))} Ω。该塔反击耐雷水平为"
            f"{format_number(ledger_record.get('backflashover_withstand_ka'))} kA，绕击耐雷水平为"
            f"{format_number(ledger_record.get('shielding_failure_withstand_ka'))} kA。"
        )
    else:
        lightning_section = "本次未匹配到可用的故障杆塔台账，杆塔结构及耐雷水平参数暂无法核实。"
    lightning_section += (
        f"本次雷电流幅值为{format_number(current)} kA。" if has_value(current) else "本次未提供可用的雷电流幅值。"
    )
    lightning_section += (
        f"根据雷电流幅值与台账耐雷水平的固定阈值比较，"
        f"{display(lightning_cause.get('report_text'), '耐雷水平数据不足')}。"
    )

    initial_phase = model.get("main_phase") or safe_get(wave, "fault_phase_review", "initial_wave_main_phase")
    active = review.get("initial_wave_active_phases") or model.get("active_phases")
    if isinstance(active, Sequence) and not isinstance(active, str):
        active = "/".join(str(item) for item in active)
    lightning_score = model.get("ln_lightning_score")
    non_lightning_score = model.get("ln_non_lightning_score")
    score_text = (
        f"模型判断雷击概率{format_number(float(lightning_score) * 100, 2)}%，"
        f"非雷击概率{format_number(float(non_lightning_score) * 100, 2)}%"
        if has_value(lightning_score) and has_value(non_lightning_score) else "雷击与非雷击概率未提供"
    )
    wave_initial_result = (
        f"{display(wave.get('initial_classification'))}；行波判别故障相{display(initial_phase)}、"
        f"候补故障相{display(active)}；{score_text}；{wave_fr_probability_sentence(model)}"
    )
    fr_probabilities = (
        (model.get("fr_backflash_probability"), model.get("fr_shielding_probability"))
        if has_value(model.get("fr_backflash_probability")) and has_value(model.get("fr_shielding_probability"))
        else normalized_two_class_probabilities(model.get("fr_backflash_score"), model.get("fr_shielding_score"))
    )
    if has_value(lightning_score) and has_value(non_lightning_score):
        probability_summary = (
            f"模型判断雷击概率{format_number(float(lightning_score) * 100, 2)}%、"
            f"非雷击概率{format_number(float(non_lightning_score) * 100, 2)}%；"
        )
    else:
        probability_summary = "雷击与非雷击模型概率未提供；"
    if fr_probabilities and all(has_value(item) for item in fr_probabilities):
        probability_summary += (
            f"模型判断反击概率{format_number(float(fr_probabilities[0]) * 100, 2)}%、"
            f"绕击概率{format_number(float(fr_probabilities[1]) * 100, 2)}%。"
        )
    else:
        probability_summary += "反击与绕击模型概率未提供。"
    wave_formats = "/".join(sorted({
        Path(str(item.get("file_name") or "")).suffix.lstrip(".").upper()
        for item in inputs.get("traveling_wave_files", []) if item.get("file_name")
    })) or "未提供"
    wave_input = (
        f"初判：{display(wave.get('initial_classification'))}；初判相：{display(initial_phase)}；"
        f"幅值信息：{translate_wave_text(review.get('amplitude_band'))}；格式：{wave_formats}"
    )
    optical_input = (
        f"位置：{display(optical.get('location') or optical.get('device_tower'))}；"
        f"现象/结论：{display(optical.get('result') or optical.get('conclusion'))}"
        if optical else "本次未提供光学监测数据"
    )
    source_by_key = {str(item.get("key")): item for item in location.get("sources", [])}
    source_by_label = {str(item.get("label")): item for item in location.get("sources", [])}
    lightning_source = source_by_key.get("lightning") or source_by_label.get("雷电定位") or {}
    record_source = source_by_key.get("recording_distance") or source_by_label.get("录波测距") or {}
    wave_source = source_by_key.get("traveling_wave_distance") or source_by_label.get("行波测距") or {}

    zero_current = fault.get("zero_sequence_current")
    record_distance = fault.get("fault_distance_km") or lightning.get("recording_distance_km")
    reasons = [user_facing_model_terms(item).rstrip("。；") + "。" for item in recommendation.get("reasons", []) if has_value(item)]
    if has_strike:
        lightning_relation_sentence = (
            f"查询雷电定位系统，线路跳闸时刻附近存在落雷记录，雷电流幅值为"
            f"{value_with_unit(lightning.get('current_ka'), 'kA')}。{relation}"
        )
    else:
        lightning_relation_sentence = "查询雷电定位系统，线路跳闸时刻附近无可用落雷记录。"
    optical_sentence = (
        f"光学监测结果为{display(optical.get('result') or optical.get('conclusion'))}。"
        if optical else ""
    )
    values: dict[str, Any] = {
        "report.id": report_identifier(result),
        "report.generated_at": generated_at,
        "report.generated_at_sentence": f"生成时间：{generated_at}",
        "fault.time": display(fault.get("fault_time")),
        "fault.station": fault.get("station") if usable_text(fault.get("station")) else "未提供",
        "fault.line_name": display(line_name),
        "fault.phase_display": phase,
        "fault.reclose": display(fault.get("reclosing")),
        "summary.event_sentence": event_sentence,
        "summary.sources_sentence": sources_sentence,
        "summary.location_sentence": summary_location,
        "summary.cause_sentence": summary_cause,
        "data_sources.actual_display": "、".join(actual_sources(result)),
        "grounding.section_sentence": grounding_sentence,
        "grounding.user_conclusion": display(ground.get("conclusion"), "待复核"),
        "grounding.current_observation": current_observation,
        "grounding.voltage_observation": voltage_observation,
        "grounding.zero_sequence_observation": zero_observation,
        "resistance.category_display": resistance_category,
        "resistance.conditional_sentence": resistance_sentence(result),
        "location.summary": display(location.get("result_text")),
        "location.source_summary": source_labels,
        "location.section_sentence": location_section,
        "location.lightning_result": "位置记录可用" if lightning_source.get("available") else "未提供",
        "location.record_result": "位置记录可用" if record_source.get("available") else "未提供",
        "location.wave_result": "位置记录可用" if wave_source.get("available") else "未提供",
        "location.record_object": display(record_source.get("tower")),
        "location.wave_object": display(wave_source.get("tower")),
        "lightning.section_sentence": lightning_section,
        "lightning.threshold_statement": lightning_location_conclusion(result),
        "lightning.record_and_relation_sentence": lightning_relation_sentence,
        "cause.ledger_withstand_sentence": display(lightning_cause.get("report_text"), "耐雷水平数据不足").rstrip("。；"),
        "wave.process_summary_sentence": wave_process_summary_sentence(result),
        "wave.initial_result_sentence": wave_initial_result,
        "wave.initial_basis_sentence": wave_ln_basis_sentence(wave, model) + wave_fr_basis_sentence(wave, model),
        "wave.probability_summary_sentence": probability_summary,
        "wave.candidate_types_display": wave_candidate_types_display(model.get("candidate_types")),
        "wave.review_result_sentence": wave_review_result_sentence(wave, model, review),
        "wave.final_fault_sentence": wave_fault_conclusion(result),
        "wave.initial_review_sentence": wave_initial_review_sentence(result),
        "wave.review_summary_sentence": wave_review_summary_sentence(result),
        "wave.initial_fusion_sentence": f"行波初判为{display(wave.get('initial_classification'))}（{wave_initial_type_detail(result)}）",
        "cause.recording_phase_sentence": recording_phase_sentence(result),
        "cause.final_fault_sentence": fused_fault_conclusion(result),
        "review.status": effective_review_decision(result),
        "review.reliability_sentence": effective_review_summary(result),
        "review.follow_up_sentence": f"3）后续意见：{effective_review_decision(result)}。",
        "review.summary_sentence": f"{effective_review_decision(result)}：{effective_review_summary(result)}。",
        "review.advice": effective_review_decision(result),
        "review.reasons_multiline": "\n".join(reasons),
        "record.time": display(fault.get("fault_time")),
        "record.line_name": display(line_name),
        "record.fault_phase": normalize_fault_phase_code(fault.get("fault_phase")) or "未提供",
        "record.duration_ms": value_with_unit(fault.get("duration_ms"), "ms"),
        "record.reclose": display(fault.get("reclosing")),
        "record.fault_voltage": display(fault.get("phase_voltage")),
        "record.fault_current": display(fault.get("phase_current")),
        "record.zero_current_raw": f"{display(zero_current)}（按3I₀/I₀双口径核对）",
        "record.zero_current_max": display(zero_current),
        "record.k0": display(fault.get("zero_sequence_compensation_k")),
        "record.distance_km": value_with_unit(record_distance, "km"),
        "record.source_file": display(paths.get("fault")),
        "lightning.amplitude_ka": value_with_unit(lightning.get("current_ka"), "kA"),
        "lightning.time": display(lightning.get("location_time")),
        "lightning.nearest_tower": display(lightning.get("nearest_tower")),
        "lightning.record_distance_km": value_with_unit(lightning.get("recording_distance_km"), "km"),
        "lightning.wave_distance_km": value_with_unit(lightning.get("traveling_wave_distance_km"), "km"),
        "lightning.location_distance_km": value_with_unit(lightning.get("distance_to_line_km"), "km"),
        "lightning.source_file": display(paths.get("lightning")),
        "lightning.data_status": "已提供" if lightning else "未提供",
        "ledger.matched_object": f"{display(ledger_record.get('line_name'))} / {display(ledger_record.get('tower_no'))}",
        "ledger.tower_no": display(ledger_record.get("tower_no")),
        "ledger.tower_shape": display(ledger_record.get("tower_shape")),
        "ledger.terrain": display(ledger_record.get("terrain")),
        "ledger.altitude_m": value_with_unit(ledger_record.get("elevation_m"), "m"),
        "ledger.insulator_length": value_with_unit(ledger_record.get("insulator_length_m"), "m"),
        "ledger.grounding_resistance_ohm": value_with_unit(ledger_record.get("ground_resistance_ohm"), "Ω"),
        "ledger.shielding_level_ka": value_with_unit(ledger_record.get("shielding_failure_withstand_ka"), "kA"),
        "ledger.backflash_level_ka": value_with_unit(ledger_record.get("backflashover_withstand_ka"), "kA"),
        "ledger.risk_level": display(ledger_record.get("lightning_risk_level")),
        "ledger.source_file": display(paths.get("ledger")),
        "ledger.match_status": "已匹配" if ledger.get("matched") else "未匹配",
        "ledger.match_note": display(ledger.get("diagnostic") or ledger.get("reason") or "线路及杆塔匹配成功"),
        "wave.input_summary": wave_input,
        "wave.data_status": "已提供" if inputs.get("traveling_wave_files") else "未提供",
        "optical.input_summary": optical_input,
        "optical.data_status": "已提供" if optical else "未提供",
        "optical.conditional_sentence": optical_sentence,
    }
    return {key: user_facing_model_terms(display(value, "")) for key, value in values.items()}


def generate_report(
    data: Mapping[str, Any],
    output_path: Path,
    template_path: Path | None = None,
    front_only: bool = False,
) -> Path:
    # 兼容两种入参：完整 result_json（含 diagnosis + sections 富结构）或裸判别结果 diagnosis。
    # 报告表格优先直接消费 result_json.sections 中的 table 内容项，保证页面与文档同源渲染；
    # 旧记录无 table sections 时回退到各 build_* 的独立拼表逻辑。
    if isinstance(data, Mapping) and "diagnosis" in data:
        result_json = data
        result = result_json.get("diagnosis", {}) or {}
        sections: Sequence[Mapping[str, Any]] = result_json.get("sections") or []
    else:
        result_json = None
        result = data or {}
        sections = []
    # 模板路径解析：None 或不存在时回退内置样式
    use_template = template_path is not None and Path(template_path).exists()

    if use_template:
        from .fillable_report_template import fill_template, unresolved_fields

        template_doc = Document(template_path)
        template_text = "\n".join(paragraph.text for paragraph in template_doc.paragraphs)
        if re.search(r"\[[a-z][a-z0-9_.]*\]", template_text):
            output = fill_template(template_path, output_path, report_field_values(result))
            unresolved = unresolved_fields(output)
            if unresolved:
                raise ValueError(f"报告存在未填充字段: {', '.join(unresolved)}")
            doc = Document(output)
            sanitize_report_model_terms(doc)
            normalize_document_fonts(doc)
            doc.save(output)
            return output
        # 模板存在但无占位字段：清空正文并基于模板样式重建
        doc = Document(template_path)
        clear_document_body(doc)
    else:
        # 无模板：创建全新文档，使用内置默认样式
        doc = Document()
    build_styles(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("输电线路故障综合判别报告")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(DARK)
    subtitle_text = "第一轮前半部分审阅稿｜" if front_only else ""
    subtitle = doc.add_paragraph(f"{subtitle_text}生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "1  报告总览")
    build_basic_information(doc, result, front_only, sections)
    build_executive_summary(doc, result)
    build_ground_and_resistance_section(doc, result)
    if not front_only:
        build_location_section(doc, result, sections)
        build_lightning_section(doc, result)
        add_heading(doc, "5  分布式行波分析与故障原因融合")
        build_wave_section(doc, result, sections)
        build_fusion_section(doc, result, sections)
        build_recommendation_section(doc, result)
        build_input_appendix(doc, result, sections)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "输电线路故障综合判别报告"
    doc.core_properties.author = "故障综合判别系统"
    sanitize_report_model_terms(doc)
    normalize_document_fonts(doc)
    doc.save(output_path)
    return output_path.resolve()


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="根据步骤1-5判别结果生成正式Word报告")
    parser.add_argument("--input-json", required=True, type=Path, help="步骤1-5判别结果JSON")
    parser.add_argument("--output-docx", required=True, type=Path, help="输出Word报告路径")
    parser.add_argument("--template-docx", type=Path, default=DEFAULT_TEMPLATE, help="Word模板路径")
    parser.add_argument("--front-only", action="store_true", help="仅生成标题、综合摘要、接地及过渡电阻部分")
    args = parser.parse_args(argv)
    if not args.input_json.exists():
        raise FileNotFoundError(f"判别结果不存在: {args.input_json}")
    result = json.loads(args.input_json.read_text(encoding="utf-8-sig"))
    output = generate_report(result, args.output_docx, args.template_docx, args.front_only)
    print(f"Word报告已生成: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
