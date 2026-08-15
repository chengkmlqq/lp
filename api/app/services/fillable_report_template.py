# NOTE(LP-343): 移植自交付版 fillable_report_template.py，用于可填充模板的字段替换。
# 仅保留 fill_template/unresolved_fields 及依赖，build_fillable_template 已移除。
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.oxml.ns import qn

FIELD_RE = re.compile(r"\[([a-z][a-z0-9_.]*)\]")


def normalize_document_fonts(doc: Document) -> None:
    for style in doc.styles:
        try:
            style.font.name = "Microsoft YaHei"
            r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
            r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
            r_fonts.set(qn("w:eastAsia"), "微软雅黑")
            r_fonts.set(qn("w:cs"), "Microsoft YaHei")
        except (AttributeError, TypeError):
            continue
    for paragraph in _all_paragraphs(doc):
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
            r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
            r_fonts.set(qn("w:eastAsia"), "微软雅黑")
            r_fonts.set(qn("w:cs"), "Microsoft YaHei")


def _all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _remove_catalog_from_report(doc: Document) -> None:
    body = doc._element.body
    start = None
    for child in list(body):
        texts = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
        if texts.startswith("附表B  占位字段完整输出字典"):
            start = child
            break
    if start is None:
        return
    remove = False
    for child in list(body):
        if child is start:
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def fill_template(template_path: Path, output_path: Path, fields: Mapping[str, Any]) -> Path:
    doc = Document(template_path)
    _remove_catalog_from_report(doc)
    values = {key: "" if value is None else str(value) for key, value in fields.items()}
    for paragraph in _all_paragraphs(doc):
        text = paragraph.text
        if not FIELD_RE.search(text):
            continue
        replaced = FIELD_RE.sub(lambda match: values.get(match.group(1), f"[{match.group(1)}]"), text)
        _set_paragraph_text(paragraph, replaced)
    normalize_document_fonts(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "输电线路故障综合判别报告"
    doc.core_properties.author = "故障综合判别系统"
    doc.save(output_path)
    return output_path.resolve()


def unresolved_fields(path: Path) -> list[str]:
    doc = Document(path)
    found: set[str] = set()
    for paragraph in _all_paragraphs(doc):
        found.update(FIELD_RE.findall(paragraph.text))
    return sorted(found)
