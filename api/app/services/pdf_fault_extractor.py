from __future__ import annotations

import csv
import io
import os
import tempfile
from pathlib import Path
from typing import Any

from fastapi import HTTPException, status
from pydantic import BaseModel

from ..services.fault_diagnosis_core import FAULT_FIELD_ALIASES, _has_value, canonicalize_fault_fields
from ..core.config import get_settings


class ExtractWavePdfResponse(BaseModel):
    ok: bool
    extracted: dict[str, Any]
    missing: list[str]
    ocr_used: bool
    raw_text_preview: str | None = None


_EXTRACTOR_PATH = Path(__file__).resolve().parents[2] / "录波" / "extract_fault_report.py"


_WAVE_FIELD_KEYS = [
    "station",
    "fault_time",
    "line_name",
    "fault_phase",
    "duration_ms",
    "reclosing",
    "phase_voltage",
    "phase_current",
    "zero_sequence_current",
    "zero_sequence_compensation_k",
    "line_length_km",
    "sensitive_angle_deg",
    "positive_sequence_impedance_ohm",
]


def _extract_from_bytes(content: bytes) -> tuple[dict[str, Any], bool, str | None]:
    if _EXTRACTOR_PATH.exists() is False:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="录波 PDF 提取器不存在，请联系管理员",
        )

    import importlib.util

    spec = importlib.util.spec_from_file_location("fault_report_extractor", str(_EXTRACTOR_PATH))
    if spec is None or spec.loader is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="加载录波 PDF 提取器失败，请联系管理员",
        )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    tmp_pdf: str | None = None
    try:
        fd, tmp_pdf = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        Path(tmp_pdf).write_bytes(content)

        text = module.extract_pdf_text(Path(tmp_pdf))
        ocr_used = False
        if not text:
            text = module.ocr_pdf(
                Path(tmp_pdf),
                engine="auto",
                dpi=150,
                tesseract_lang="chi_sim+eng",
                pages=None,
            )
            ocr_used = True

        if not text:
            return {}, False, None

        preview = text[:200]
        parsed = dict(module.parse_fault_info(text))
        return parsed, ocr_used, preview
    finally:
        if tmp_pdf:
            try:
                Path(tmp_pdf).unlink(missing_ok=True)
            except Exception:
                pass


def _build_wave_response(parsed: dict[str, Any], ocr_used: bool, preview: str | None) -> ExtractWavePdfResponse:
    canonical = canonicalize_fault_fields(parsed)
    extracted: dict[str, Any] = {key: canonical.get(key) for key in _WAVE_FIELD_KEYS}
    missing = [key for key in _WAVE_FIELD_KEYS if not extracted.get(key)]
    return ExtractWavePdfResponse(
        ok=len(missing) == 0,
        extracted=extracted,
        missing=missing,
        ocr_used=ocr_used,
        raw_text_preview=preview,
    )


def _all_aliases() -> set[str]:
    return {alias for aliases in FAULT_FIELD_ALIASES.values() for alias in aliases}


def _read_tabular_content(content: bytes, ext: str) -> dict[str, Any]:
    """读取 txt/csv/xlsx 表格内容并提取故障字段键值：

    1) 横向表头：在前 12 行中找故障字段别名命中最多的一行作为表头，取其下首个有效数据行
       （容忍标题行/空行/合并表头前的说明行）；
    2) 纵向布局：横向未命中时，按每行 [标签, 值] 两列解析。
    """
    raw_rows: list[list[Any]] = []
    if ext == "xlsx":
        fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            Path(tmp_path).write_bytes(content)
            from openpyxl import load_workbook  # type: ignore

            wb = load_workbook(tmp_path, read_only=True, data_only=True)
            try:
                sheet = wb[wb.sheetnames[0]]
                for values in sheet.iter_rows(values_only=True):
                    if any(_has_value(v) for v in values):
                        raw_rows.append(list(values))
            finally:
                wb.close()
        finally:
            try:
                Path(tmp_path).unlink(missing_ok=True)
            except Exception:
                pass
    else:  # txt / csv
        text = content.decode("utf-8-sig")
        raw_rows = [list(csv.reader([line]))[0] for line in text.splitlines() if line.strip()]

    if not raw_rows:
        return {}

    aliases = _all_aliases()

    # 1) 横向表头定位（至少命中 2 个别名，避免把纵向 label/value 布局误判为横向表头）
    best_idx, best_score = -1, 0
    for i, row in enumerate(raw_rows[:12]):
        score = sum(1 for cell in row if str(cell).strip() in aliases)
        if score > best_score:
            best_idx, best_score = i, score
    if best_idx >= 0 and best_score >= 2:
        headers = [str(c).strip() for c in raw_rows[best_idx]]
        for row in raw_rows[best_idx + 1:]:
            data = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
            if any(_has_value(v) for v in data.values()):
                return data

    # 2) 纵向布局：[标签, 值] 两列
    vertical: dict[str, Any] = {}
    for row in raw_rows:
        if len(row) >= 2 and str(row[0]).strip() in aliases:
            vertical[str(row[0]).strip()] = row[1]
    return vertical


def extract_wave_file(content: bytes, file_name: str = "upload") -> ExtractWavePdfResponse:
    """按扩展名分发提取 13 个故障字段：pdf→文本/OCR；txt/csv/xlsx→表格解析（横向表头或纵向 label/value）；
    json→对象/数组首元素。统一经 canonicalize_fault_fields 归一化后返回，不回写数据库。"""
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="上传文件为空")

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext == "pdf":
        parsed, ocr_used, preview = _extract_from_bytes(content)
        return _build_wave_response(parsed, ocr_used, preview)

    if ext in {"txt", "csv", "json", "xlsx"}:
        if ext == "json":
            import json as _json

            try:
                parsed = _json.loads(content.decode("utf-8-sig"))
            except Exception as exc:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"JSON 解析失败：{exc}") from exc
            if isinstance(parsed, list):
                parsed = parsed[0] if parsed else {}
            if not isinstance(parsed, dict):
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="JSON 顶层应为对象或数组")
            return _build_wave_response(parsed, False, str(parsed)[:200])

        try:
            parsed = _read_tabular_content(content, ext)
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"{ext.upper()} 解析失败：{exc}") from exc
        if not parsed:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="未识别到故障字段：请使用「下载模板」的格式（录波记录厂站/故障时刻/故障线路名称等表头）或上传录波故障报告 PDF",
            )
        return _build_wave_response(parsed, False, str(parsed)[:200])

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="暂不支持该格式，录波数据支持 pdf/txt/csv/json/xlsx",
    )
