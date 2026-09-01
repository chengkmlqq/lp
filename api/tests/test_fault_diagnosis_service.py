from __future__ import annotations

from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import get_settings
from app.core.database import Base
from app.models.fault_analysis import FaultAnalysis
from app.models.user import User
from app.schemas.fault_analysis import FaultAnalysisCreateRequest
from app.services import fault_analysis_service
from app.services import fault_diagnosis_service


def _build_sessionmaker(tmp_path: Path) -> sessionmaker:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    return sessionmaker(bind=engine, autocommit=False, autoflush=False, expire_on_commit=False)


def _sample_user(session: Session) -> User:
    user = User(id="u1", username="tester", status="enabled", password_hash="x")
    session.add(user)
    session.commit()
    return user


def _sample_analysis(session: Session) -> FaultAnalysis:
    item = FaultAnalysis(
        id="fa-sample",
        title="龙和线故障",
        station="广东河源500kV上寨变",
        line_name="220kV 龙和线",
        fault_time="2023-03-27 20:20:31.939",
        fault_phase="AN",
        fault_duration_ms=64,
        reclosure_status="success",
        fault_voltage_v=37.04,
        fault_current_a=1.3,
        max_zero_current_a=1.45,
        zero_compensation=0.52,
        line_length_km=89.77,
        line_angle_deg=85,
        line_impedance_ohm=28.520,
        lightning_current_ka=20.3,
        wave_distance_km=16.022,
        traveling_wave_distance_km=16.51,
        status="pending",
        result_json={},
        upload_files={},
    )
    session.add(item)
    session.commit()
    return item


@pytest.fixture()
def db_session(tmp_path, monkeypatch):
    SessionLocal = _build_sessionmaker(tmp_path)
    session = SessionLocal()
    # 将判别服务的共享卷指向临时目录，避免写入仓库 data/。
    settings = get_settings()
    shim = _SettingsShim(
        str(tmp_path / "vfs"),
        settings.fault_diag_ledger_csv,
        settings.fault_diag_matlab_compiled_dir,
    )
    monkeypatch.setattr(fault_diagnosis_service, "get_settings", lambda: shim)
    monkeypatch.setattr(fault_analysis_service, "get_settings", lambda: shim)
    stored: dict[str, bytes] = {}

    def fake_put_object(key: str, data: bytes, content_type: str | None = None) -> str:
        stored[key] = data
        return key

    def fake_download_object_to_file(key: str, target: Path) -> Path:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(stored[key])
        return target

    monkeypatch.setattr(fault_diagnosis_service, "put_object", fake_put_object)
    monkeypatch.setattr(fault_diagnosis_service, "download_object_to_file", fake_download_object_to_file)
    monkeypatch.setattr(fault_analysis_service, "put_object", fake_put_object)
    session._test_minio_objects = stored
    yield session
    session.close()


class _SettingsShim:
    def __init__(self, vfs_root: str, ledger_csv: str, compiled_dir: str) -> None:
        self.file_vfs_root = vfs_root
        self.fault_diag_ledger_csv = ledger_csv
        self.fault_diag_matlab_compiled_dir = compiled_dir


def test_run_fault_diagnosis_completes_and_backfills(db_session: Session):
    _sample_user(db_session)
    item = _sample_analysis(db_session)

    result = fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)

    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, item.id)
    assert updated.status == "completed"
    assert updated.risk_level in {"low", "medium", "high"}
    assert updated.fault_type  # 例如 "A相单相接地"
    assert updated.fault_phase == "A"
    assert updated.confidence is not None
    assert updated.confidence > 0

    assert isinstance(updated.result_json, dict)
    assert updated.result_json["artifacts"]["result_json"] == "fault-analysis/fa-sample/Output/result.json"
    assert "sections" in updated.result_json
    assert updated.result_json["sections"]
    assert "diagnosis" in updated.result_json
    assert updated.result_json["diagnosis"]["ground_fault_decision"]["is_ground_fault"] is True
    assert result == updated.result_json


def test_run_fault_diagnosis_degrades_with_minimal_input(db_session: Session):
    _sample_user(db_session)
    item = FaultAnalysis(
        id="fa-minimal",
        title="最小输入",
        station="某站",
        fault_time="2023-01-01 10:00:00",
        line_name="某线路",
        fault_phase="B",
        status="pending",
        result_json={},
        upload_files={},
    )
    db_session.add(item)
    db_session.commit()

    result = fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)

    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, item.id)
    # 判别系统在输入不足时不应抛异常，而是降级输出。
    assert updated.status == "completed"
    assert "sections" in updated.result_json
    assert updated.fault_type is not None


def test_run_fault_diagnosis_with_traveling_wave_file(db_session: Session, tmp_path: Path):
    _sample_user(db_session)
    # 构造一个可解析的行波 TXT（Key=Value + CurrentDataInput）。
    samples = ",".join(str(value) for value in [1.0] * 40)
    wave_dir = tmp_path / "vfs" / "fault_analysis" / "fa-wave"
    wave_dir.mkdir(parents=True)
    wave_file = wave_dir / "wave_data_file_A相_1号杆塔.txt"
    wave_file.write_text(f"SampleFrequency=1000000\nCurrentDataInput={samples}\n", encoding="utf-8")

    item = FaultAnalysis(
        id="fa-wave",
        title="行波测试",
        station="某站",
        fault_time="2023-01-01 10:00:00",
        line_name="某线路",
        fault_phase="AN",
        fault_duration_ms=64,
        fault_voltage_v=37.04,
        fault_current_a=1.3,
        max_zero_current_a=1.45,
        zero_compensation=0.52,
        line_length_km=89.77,
        line_angle_deg=85,
        line_impedance_ohm=28.520,
        status="pending",
        result_json={},
        upload_files={"traveling_wave_file": wave_file.relative_to(tmp_path / "vfs").as_posix()},
    )
    db_session.add(item)
    db_session.commit()

    result = fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)

    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, item.id)
    assert updated.status == "completed"
    wave_analysis = updated.result_json["diagnosis"]["traveling_wave_analysis"]
    assert wave_analysis["phase_analyses"]  # 已解析行波相
    assert wave_analysis["status"] in {"analyzed", "needs_review"}


def test_store_upload_file_and_create_analysis(db_session: Session, tmp_path: Path):
    _sample_user(db_session)
    metadata = fault_analysis_service.store_analysis_upload_file(
        "fa-upload", "wave_data_file", "a.txt", b"hello"
    )
    assert metadata == {
        "key": "fault-analysis/fa-upload/Input/wave_data_file/a.txt",
        "filename": "a.txt",
        "content_type": "application/octet-stream",
        "source": "uploaded",
    }
    assert db_session._test_minio_objects[metadata["key"]] == b"hello"


def test_run_fault_diagnosis_materializes_minio_input(db_session: Session):
    _sample_user(db_session)
    key = "fault-analysis/fa-minio/Input/traveling_wave_file/A相.txt"
    db_session._test_minio_objects[key] = b"SampleFrequency=1000000\nCurrentDataInput=1,1,1,1,1,1,1,1\n"
    item = FaultAnalysis(
        id="fa-minio",
        title="MinIO 行波测试",
        station="某站",
        fault_time="2023-01-01 10:00:00",
        line_name="某线路",
        fault_phase="AN",
        fault_voltage_v=37.04,
        fault_current_a=1.3,
        max_zero_current_a=1.45,
        status="pending",
        result_json={},
        upload_files={
            "traveling_wave_file": [
                {
                    "key": key,
                    "filename": "A相.txt",
                    "content_type": "text/plain",
                    "source": "uploaded",
                }
            ]
        },
    )
    db_session.add(item)
    db_session.commit()

    fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)

    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, item.id)
    assert updated.status == "completed"
    assert updated.result_json["upload_files"]["traveling_wave_file"][0]["key"] == key
    assert updated.result_json["artifacts"]["result_json"] == "fault-analysis/fa-minio/Output/result.json"


def test_create_analysis_dispatches_and_fails_gracefully_without_broker(
    db_session: Session, monkeypatch
):
    _sample_user(db_session)
    from app.tasks import fault_diagnosis_tasks

    def _raise(*args, **kwargs):
        raise ConnectionError("broker unavailable")

    monkeypatch.setattr(fault_diagnosis_tasks.run_fault_diagnosis_task, "delay", _raise)
    payload = FaultAnalysisCreateRequest(
        title="创建测试",
        station="某站",
        fault_time="2023-01-01 10:00:00",
        line_name="某线路",
        fault_phase="AN",
        line_id="l1",
    )

    user = db_session.query(User).first()
    response = fault_analysis_service.create_analysis(
        db_session, payload, actor=user, auto_diagnose=True
    )
    assert response.analysis.id
    # 队列不可用时创建仍成功，记录被标记为失败并记录原因。
    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, response.analysis.id)
    assert updated.status == "failed"
    assert "error" in updated.result_json


def test_sections_contain_structured_tables(db_session: Session):
    """sections 富结构包含 docx 对应表格（type:"table"），章节标题与报告对齐。"""
    _sample_user(db_session)
    item = _sample_analysis(db_session)

    result = fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)
    sections = result["sections"]

    # 章节标题与 docx 报告对齐
    titles = [s["title"] for s in sections]
    assert "1 报告总览" in titles
    assert "2 接地故障及过渡电阻判别" in titles
    assert "3 故障定位" in titles
    assert "4 耐雷水平与雷击原因分析" in titles
    assert "5.1 分布式行波分析" in titles
    assert "5.2 故障原因融合" in titles
    assert "6 总体状态与后续意见" in titles
    assert "A.1 录波故障信息" in titles
    assert "A.2 雷电流及定位信息" in titles
    assert "A.3 线路台账匹配信息" in titles
    assert "A.4 行波与光学监测信息" in titles

    tables = [
        item
        for section in sections
        for item in section["content"]
        if isinstance(item, dict) and item.get("type") == "table"
    ]
    # 至少包含 7 张结构化表（表2 定位信息来源仅在存在可用定位源时出现）
    assert len(tables) >= 7

    overview = next((t for t in tables if t.get("title", "").startswith("表1")), None)
    assert overview is not None
    assert overview["layout"] == "key_value"
    # 2×2 键值对布局 → 4 列、不少于 5 行
    assert len(overview["columns"]) == 4
    assert len(overview["rows"]) >= 5


def test_generate_report_docx_consumes_section_tables(db_session: Session, tmp_path: Path):
    """docx 报告生成优先消费 result_json.sections 中的 table，保证与页面同源。"""
    item = _sample_analysis(db_session)
    fault_diagnosis_service.run_fault_diagnosis(db_session, item.id)
    db_session.expire_all()
    updated = db_session.get(FaultAnalysis, item.id)

    from app.services import fault_diagnosis_report
    from docx import Document

    output_path = tmp_path / "report.docx"
    fault_diagnosis_report.generate_report(
        updated.result_json, output_path, template_path=None, front_only=False
    )
    assert output_path.is_file()

    doc = Document(str(output_path))
    section_tables = [
        content
        for section in updated.result_json["sections"]
        for content in section.get("content", [])
        if isinstance(content, dict) and content.get("type") == "table"
    ]
    # 表1 内容与 docx 第一个表格（故障分析总览）逐项一致（同一份数据源）
    if section_tables:
        overview = next(
            (t for t in section_tables if t.get("title", "").startswith("表1")), None
        )
        if overview is not None and doc.tables:
            doc_overview = doc.tables[0]
            for idx, cell in enumerate(doc_overview.rows[1].cells):
                assert str(overview["rows"][0][idx]) == cell.text


def test_finalize_artifacts_archives_matlab_directory_and_debug_inputs(
    db_session: Session, tmp_path: Path
):
    """MATLAB 全量产物目录递归归档 + 判别中间输入 debug JSON 落盘，清单写入 artifacts。"""
    _sample_user(db_session)
    item = _sample_analysis(db_session)

    run_output_dir = tmp_path / "module_runs" / "20260817_000000_000001"
    run_output_dir.mkdir(parents=True)
    (run_output_dir / "classification_result.mat").write_bytes(b"\x00MATLAB-v7\x01")
    (run_output_dir / "classification_summary.xlsx").write_bytes(b"PK\x03\x04")
    prediction_file = run_output_dir / "wave_module_prediction.txt"
    prediction_file.write_text("mechanism=雷击\n", encoding="utf-8")
    nested = run_output_dir / "sub" / "wave_input_file_list.xlsx"
    nested.parent.mkdir(parents=True)
    nested.write_bytes(b"PK\x03\x04")
    (run_output_dir / "占位.txt").write_text("hello", encoding="utf-8")

    result_json = fault_diagnosis_service._finalize_artifacts(
        {"sections": [], "diagnosis": {}},
        analysis_id=item.id,
        wave_model_path=prediction_file,
        matlab_output_dir=run_output_dir,
        debug_inputs=(
            {"station": "某站"},
            {"query_result": "否"},
            [{"line_name": "某线路"}],
            [{"path": "/tmp/wave.txt"}],
        ),
    )

    artifacts = result_json["artifacts"]
    assert artifacts["result_json"] == f"fault-analysis/{item.id}/Output/result.json"
    assert (
        artifacts["matlab_wave_prediction"]
        == f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/wave_module_prediction.txt"
    )
    # MATLAB 全量产物：保持 run_id + 嵌套相对路径
    assert artifacts["matlab_outputs"] == [
        f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/classification_result.mat",
        f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/classification_summary.xlsx",
        f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/sub/wave_input_file_list.xlsx",
        f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/wave_module_prediction.txt",
        f"fault-analysis/{item.id}/Output/matlab/20260817_000000_000001/占位.txt",
    ]
    # debug 输入 JSON 落盘清单
    assert artifacts["debug_outputs"] == [
        f"fault-analysis/{item.id}/Output/debug/fault_raw.json",
        f"fault-analysis/{item.id}/Output/debug/lightning_raw.json",
        f"fault-analysis/{item.id}/Output/debug/ledger_rows.json",
        f"fault-analysis/{item.id}/Output/debug/traveling_wave_records.json",
    ]

    minio_objects = db_session._test_minio_objects
    assert minio_objects[artifacts["matlab_outputs"][2]] == b"PK\x03\x04"
    assert "某站".encode("utf-8") in minio_objects[artifacts["debug_outputs"][0]]


def test_finalize_artifacts_no_matlab_keeps_compat_artifacts(db_session: Session):
    """未启用 MATLAB 引擎路径时，artifacts 保留 result_json，且 matlab_outputs 不存在。"""
    _sample_user(db_session)
    item = _sample_analysis(db_session)

    result_json = fault_diagnosis_service._finalize_artifacts(
        {"sections": [], "diagnosis": {}},
        analysis_id=item.id,
        wave_model_path=None,
        matlab_output_dir=None,
        debug_inputs=None,
    )

    artifacts = result_json["artifacts"]
    assert artifacts == {"result_json": f"fault-analysis/{item.id}/Output/result.json"}
    assert "matlab_outputs" not in artifacts
    assert "debug_outputs" not in artifacts



def test_list_analysis_summary_includes_reclosure_status(db_session: Session):
    """列表 summary 序列化必须携带重合闸情况（此前 serialize_summary 漏掉该字段导致列表展示为空）。"""
    _sample_user(db_session)
    _sample_analysis(db_session)

    result = fault_analysis_service.list_analyses(db_session, keyword="龙和线")
    assert result.total == 1
    item = result.items[0]
    assert item.reclosure_status == "success"
    assert item.id == "fa-sample"
